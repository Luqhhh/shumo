"""Read-only Q4 design checks; no formal predictor, optimizer or case execution."""

from __future__ import annotations

import argparse
import datetime as dt
import json
import math
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from microgrid.dataio import read_wide_attachment, sha256_file


def audit(repo: Path) -> dict:
    price_path = repo / "data/raw/附件4.xlsx"
    rows = read_wide_attachment(price_path, sheet_name="Sheet1", kind="price_actual", unit="元/kWh")
    times = [dt.datetime.fromisoformat(row["parsed_timestamp"]) for row in rows]
    prices = [float(row["value"]) for row in rows]
    ordered = sorted(times)
    expected = [dt.datetime(2025, 1, 1) + dt.timedelta(minutes=10 * i) for i in range(1, 52561)]
    if ordered != expected or len(set(times)) != len(times):
        raise ValueError("prices must cover all 52560 unique 2025 ten-minute endpoints")
    if any(not math.isfinite(value) or value < 0 for value in prices):
        raise ValueError("non-finite or negative actual price: stop; do not clip input")
    templates = {}
    for name in ("result4-2.xlsx", "result4-3.xlsx"):
        path = repo / "data/templates" / name
        wb = load_workbook(path, data_only=False, read_only=True)
        try:
            sheets = {}
            for ws in wb:
                sheets[ws.title] = {
                    "rows": ws.max_row,
                    "columns": ws.max_column,
                    "ellipsis_rows": sorted(
                        {cell.row for row in ws for cell in row if cell.value == "⁝"}
                    ),
                }
                if ws.title in ("计划购电量", "调整购电量"):
                    sheets[ws.title]["boundary_labels"] = {
                        key: str(ws[key].value) for key in ("B1", "EN1", "EO1", "EP1", "EQ1")
                    }
                    dates = [
                        row[0].date()
                        for row in ws.iter_rows(min_row=2, max_row=335, max_col=1, values_only=True)
                    ]
                    if dates != [dt.date(2025, 2, 1) + dt.timedelta(days=i) for i in range(334)]:
                        raise ValueError(f"{name}/{ws.title}: unexpected date grid")
            templates[name] = {"sha256": sha256_file(path), "sheets": sheets}
        finally:
            wb.close()
    # Algebra checks use synthetic quantities, not official operating outcomes.
    rng = np.random.default_rng(20260912)
    load, pv, g = rng.uniform(0, 2000, size=(3, 10000))
    cbar, dbar = rng.uniform(0, 5000 / 6, size=(2, 10000))
    charging = rng.integers(0, 2, size=10000).astype(bool)
    cbar[~charging] = 0
    dbar[charging] = 0
    c = np.minimum(cbar, np.maximum(g + pv - load, 0))
    d = np.minimum(dbar, load)
    u = np.maximum(load + c - g - pv - d, 0)
    pv_use = np.minimum(pv, load + c - d)
    grid_use = load + c - d - u - pv_use
    spill, curtail = g - grid_use, pv - pv_use
    residual = g + pv_use + d + u - load - c - spill
    assert np.max(np.abs(residual)) <= 1e-6
    assert all(np.min(x) >= -1e-6 for x in (c, d, u, pv_use, grid_use, spill, curtail))
    assert np.all((u <= 1e-6) | ((c <= 1e-6) & (spill <= 1e-6) & (curtail <= 1e-6)))
    assert np.all((c <= 1e-6) | (d <= 1e-6))
    bill = 10 * 2 + 0.5 * 1 * 2 + 1.5 * 3 * 1 + 5 * 2 * 1
    assert bill == 35.5
    return {
        "schema_version": 1,
        "status": "design_evidence_only",
        "formal_run": False,
        "price": {
            "sha256": sha256_file(price_path),
            "count": len(prices),
            "min": min(prices),
            "max": max(prices),
            "zero_count": prices.count(0),
            "negative_count": 0,
            "nonfinite_count": 0,
            "full_grid": True,
        },
        "templates": templates,
        "synthetic_formula_checks": {
            "seed": 20260912,
            "cases": 10000,
            "max_ledger_residual_kwh": float(np.max(np.abs(residual))),
            "feedback_invariants_pass": True,
            "bill_fixture_cny": bill,
        },
        "not_tested": [
            "forecast_accuracy",
            "milp_implementation",
            "annual_run",
            "terminal_success_under_actual_feedback",
            "xlsx_export_readback",
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo", type=Path, default=Path("."))
    parser.add_argument(
        "--output", type=Path, default=Path("outputs/evidence/q4_design_audit.json")
    )
    args = parser.parse_args()
    result = audit(args.repo)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True, allow_nan=False) + "\n"
    )
    print(args.output)


if __name__ == "__main__":
    main()
