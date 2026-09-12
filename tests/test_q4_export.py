from __future__ import annotations

import datetime as dt

import pytest
from openpyxl import load_workbook

from microgrid.problem.contracts import BatteryAction, BatteryState, CaseResult, IntervalResult
from microgrid.problem.purchase_ledger import PurchaseLedger
from microgrid.problem.q4_common import ACTION_START, STEP
from microgrid.problem.q4_export import fill_workbook


@pytest.mark.parametrize("case_id", ["q4_2", "q4_3"])
def test_synthetic_annual_template_mapping_readback_and_expansion(synthetic_template_repo, case_id):
    path = (
        synthetic_template_repo
        / "data"
        / "templates"
        / ("result4-2.xlsx" if case_id == "q4_2" else "result4-3.xlsx")
    )
    wb = load_workbook(path)
    ledger = PurchaseLedger(case_id)
    intervals = []
    state, action = BatteryState(6000), BatteryAction()
    for i in range(334):
        midnight = ACTION_START + dt.timedelta(days=i)
        wb["计划购电量"].cell(i + 2, 1, str(midnight.date()))
        ledger.submit(midnight, tuple(midnight + k * STEP for k in range(144)), (100.0,) * 144)
        if case_id == "q4_3":
            wb["调整购电量"].cell(i + 2, 1, str(midnight.date()))
            for hour, value in ((6, 101.0), (12, 102.0), (18, 103.0)):
                event = midnight + dt.timedelta(hours=hour)
                ledger.submit(
                    event,
                    tuple(event + k * STEP for k in range(144 - hour * 6)),
                    (value,) * (144 - hour * 6),
                )
        for k in range(144):
            slot = midnight + k * STEP
            final = ledger.committed(slot)
            emergency = 1.0 if k == 143 else 0.0
            ledger.settle(slot, available_at=slot + STEP, actual_price=2, emergency_kwh=emergency)
            intervals.append(
                IntervalResult(
                    midnight.date(),
                    k,
                    (final + emergency) * 6,
                    0,
                    100,
                    final,
                    emergency,
                    action,
                    state,
                    state,
                )
            )
    result = CaseResult(
        case_id, "synthetic-mapping-test", "success", True, intervals=tuple(intervals)
    )
    mappings, expected = fill_workbook(wb, result, ledger)
    output = synthetic_template_repo / f"SYNTHETIC-{case_id}.xlsx"
    wb.save(output)
    wb.close()
    readback = load_workbook(output, data_only=True)
    try:
        for sheet, row, col, value, tolerance in expected:
            actual = readback[sheet].cell(row, col).value
            assert actual == value if tolerance is None else abs(actual - value) <= tolerance
        assert len(mappings) == 48096
        assert readback["计划购电量"]["B2"].value == 100
        assert readback["计划购电量"]["EO335"].value == 100
        assert readback["计划购电量"]["EQ2"].value == 28800
        assert readback["充放电量"].max_row == 2005
        assert readback["充放电量"]["F2"].value == 6000
        assert readback["紧急购电量"].max_row == 335
        assert readback["紧急购电量"]["B2"].value == "23:50-0:00+1"
        if case_id == "q4_3":
            assert readback["调整购电量"]["B2"].value == 100
            assert readback["调整购电量"]["EO2"].value == 103
            assert readback["调整购电量"]["EQ2"].value == 648
    finally:
        readback.close()
