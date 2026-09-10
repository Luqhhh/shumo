from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "src"))


@pytest.fixture
def tiny_attachment1(tmp_path: Path) -> Path:
    path = tmp_path / "附件1.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["时间", "电价", "小区负载", "光伏发电预测功率"])
    ws.append([6.9444444444444441e-3, 0.4248, 3439.8466, 0.0])
    ws.append(["0:20", 0.4245, 3437.9792, 0.0])
    ws.append(["0:00+1", 0.4276, 3444.7259, 0.0])
    wb.save(path)
    return path


@pytest.fixture
def tiny_forecast(tmp_path: Path) -> Path:
    path = tmp_path / "附件3.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    headers = ["日期", "预报时刻"] + [f"预报{i}小时" for i in range(1, 8)]
    ws.append(headers)
    ws.append(["2025-1-1", "0:00"] + [float(i) for i in range(1, 8)])
    ws.append(["", "6:00"] + [float(10 + i) for i in range(1, 8)])
    ws.append(["", "12:00"] + [float(20 + i) for i in range(1, 8)])
    ws.append(["", "18:00"] + [float(30 + i) for i in range(1, 8)])
    ws.append(["2025-1-2", "0:00"] + [float(40 + i) for i in range(1, 8)])
    wb.save(path)
    return path


@pytest.fixture
def synthetic_template_repo(tmp_path: Path) -> Path:
    repo = tmp_path / "repo"
    (repo / "data" / "templates").mkdir(parents=True)
    for case, filename in {
        "q1": "result1.xlsx",
        "q2": "result2.xlsx",
        "q3": "result3.xlsx",
        "q4_2": "result4-2.xlsx",
        "q4_3": "result4-3.xlsx",
    }.items():
        wb = Workbook()
        if case == "q1":
            plan = wb.active
            plan.title = "计划购电量"
            plan.append(["时间段", "购电量"])
            charge = wb.create_sheet("充放电量")
            charge.append(["时间段", "充电量", "放电量", "时刻", "储电量"])
        else:
            plan = wb.active
            plan.title = "计划购电量"
            plan["A1"] = "日期\\时间"
            plan["B1"] = "0:10-0:20"
            plan.cell(row=1, column=145, value="0:00-0:10+1")
            plan.cell(row=1, column=146, value="全天购电量")
            plan.cell(row=1, column=147, value="全天购电费")
            if case in {"q3", "q4_3"}:
                adjust = wb.create_sheet("调整购电量")
                adjust["A1"] = "日期\\时间"
                adjust["B1"] = "0:10-0:20"
                adjust.cell(row=1, column=145, value="0:00-0:10+1")
                adjust.cell(row=1, column=146, value="全天购电量")
                adjust.cell(row=1, column=147, value="全天购电费")
            charge = wb.create_sheet("充放电量")
            charge.append(["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
            emergency = wb.create_sheet("紧急购电量")
            emergency.append(["日期", "购电时间段", "购电量"])
        wb.save(repo / "data" / "templates" / filename)
    return repo
