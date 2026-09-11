from __future__ import annotations

import datetime as dt
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from microgrid.dataio import sha256_file
from microgrid.problem.contracts import (
    INITIAL_SOC_KWH,
    BatteryAction,
    BatteryState,
    CaseResult,
    IntervalResult,
)
from microgrid.problem.q2_inputs import (
    ActualInterval,
    FixedPricePoint,
    historical_info_items,
    info_set_at,
    load_q2_inputs,
    load_q4_2_prices,
    require_matching_q4_2_grid,
)
from microgrid.problem.result_io import case_result_to_dict, load_case_result, save_case_result
from microgrid.schemas import InputError


def test_actual_interval_converts_ten_minute_power_to_energy() -> None:
    actual = ActualInterval(
        day=dt.date(2025, 2, 1),
        slot=0,
        start=dt.datetime(2025, 2, 1, 0, 0),
        end=dt.datetime(2025, 2, 1, 0, 10),
        load_kw=600.0,
        pv_kw=300.0,
        load_source_ref="附件2.xlsx!负荷!B2 sha256=aaaaaaaaaaaa",
        pv_source_ref="附件2.xlsx!光伏!B2 sha256=aaaaaaaaaaaa",
    )

    assert actual.load_kw == 600.0
    assert actual.pv_kw == 300.0
    assert actual.load_kwh == pytest.approx(100.0)
    assert actual.pv_kwh == pytest.approx(50.0)


@pytest.mark.parametrize("bad_value", [-1.0, float("nan"), float("inf")])
def test_value_records_reject_negative_or_non_finite_values(bad_value: float) -> None:
    with pytest.raises(InputError, match="price_cny_per_kwh"):
        FixedPricePoint(slot=0, price_cny_per_kwh=bad_value, source_ref="cell")


def test_q2_case_result_roundtrip_preserves_adapter_provenance(tmp_path: Path) -> None:
    actual = ActualInterval(
        day=dt.date(2025, 2, 1),
        slot=0,
        start=dt.datetime(2025, 2, 1, 0, 0),
        end=dt.datetime(2025, 2, 1, 0, 10),
        load_kw=600.0,
        pv_kw=300.0,
        load_source_ref="附件2.xlsx!负荷!B2 sha256=aaaaaaaaaaaa",
        pv_source_ref="附件2.xlsx!光伏!B2 sha256=aaaaaaaaaaaa",
    )
    state = BatteryState(INITIAL_SOC_KWH)
    interval = IntervalResult(
        day=actual.day,
        slot=actual.slot,
        load_kw=actual.load_kw,
        pv_kw=actual.pv_kw,
        planned_purchase_kwh=0.0,
        adjusted_purchase_kwh=0.0,
        emergency_purchase_kwh=0.0,
        action=BatteryAction(),
        state_start=state,
        state_end=state,
        source_ref=f"{actual.load_source_ref}; {actual.pv_source_ref}",
    )
    result = CaseResult(
        case_id="q2",
        run_id="synthetic-adapter-roundtrip",
        status="interface-test",
        is_synthetic=True,
        intervals=(interval,),
        metadata={"input_hashes": [["attachment2", "a" * 64]]},
    )

    path = tmp_path / "q2-result.json"
    save_case_result(path, result)
    loaded = load_case_result(path)

    assert case_result_to_dict(loaded) == case_result_to_dict(result)
    assert loaded.is_synthetic is True
    assert loaded.intervals[0].source_ref == interval.source_ref


def _right_endpoint_labels() -> list[str]:
    labels: list[str] = []
    for slot in range(143):
        minutes = (slot + 1) * 10
        labels.append(f"{minutes // 60}:{minutes % 60:02d}")
    labels.append("0:00+1")
    return labels


def _write_attachment1(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["时间", "电价", "小区负载", "光伏发电预测功率"])
    for slot, label in enumerate(_right_endpoint_labels()):
        ws.append([label, 0.4 + slot / 1000.0, 1000.0 + slot, 200.0 + slot])
    wb.save(path)
    wb.close()
    return path


def _write_wide_workbook(
    path: Path,
    *,
    sheets: tuple[str, ...],
    days: tuple[dt.date, ...],
    omitted: dict[str, set[tuple[dt.date, int]]] | None = None,
    duplicate_header_slot: int | None = None,
    bad_values: dict[tuple[str, dt.date, int], object] | None = None,
) -> Path:
    omitted = omitted or {}
    bad_values = bad_values or {}
    labels = _right_endpoint_labels()
    if duplicate_header_slot is not None:
        labels[duplicate_header_slot] = labels[duplicate_header_slot - 1]

    wb = Workbook()
    for sheet_index, sheet_name in enumerate(sheets):
        ws = wb.active if sheet_index == 0 else wb.create_sheet()
        ws.title = sheet_name
        ws.append(["日期", *labels])
        for day_index, day in enumerate(days):
            row: list[object | None] = [day]
            for slot in range(144):
                key = (day, slot)
                value: object | None = 1000.0 + 10.0 * day_index + slot
                if key in omitted.get(sheet_name, set()):
                    value = None
                value = bad_values.get((sheet_name, day, slot), value)
                row.append(value)
            ws.append(row)
    wb.save(path)
    wb.close()
    return path


def test_load_q2_inputs_accepts_excel_time_objects_in_attachment1(tmp_path: Path) -> None:
    attachment1 = _write_attachment1(tmp_path / "附件1.xlsx")
    wb = load_workbook(attachment1)
    wb["Sheet1"]["A2"] = dt.time(0, 10)
    wb.save(attachment1)
    wb.close()
    attachment2 = _write_wide_workbook(
        tmp_path / "附件2.xlsx",
        sheets=("实际负荷", "实际光伏"),
        days=(dt.date(2025, 1, 1),),
    )

    bundle = load_q2_inputs(
        attachment1_path=attachment1,
        attachment2_path=attachment2,
        load_sheet_name="实际负荷",
        pv_sheet_name="实际光伏",
    )

    assert len(bundle.fixed_prices) == 144
    assert bundle.fixed_prices[0].slot == 0


def test_load_q2_inputs_accepts_excel_time_headers_in_attachment2(tmp_path: Path) -> None:
    attachment1 = _write_attachment1(tmp_path / "附件1.xlsx")
    attachment2 = _write_wide_workbook(
        tmp_path / "附件2.xlsx",
        sheets=("实际负荷", "实际光伏"),
        days=(dt.date(2025, 1, 1),),
    )
    wb = load_workbook(attachment2)
    wb["实际负荷"]["B1"] = dt.time(0, 10)
    wb["实际光伏"]["B1"] = dt.time(0, 10)
    wb.save(attachment2)
    wb.close()

    bundle = load_q2_inputs(
        attachment1_path=attachment1,
        attachment2_path=attachment2,
        load_sheet_name="实际负荷",
        pv_sheet_name="实际光伏",
    )

    assert len(bundle.actuals) == 144
    assert bundle.actuals[0].start == dt.datetime(2025, 1, 1, 0, 0)


def test_load_q2_inputs_aligns_fixed_prices_and_actuals(tmp_path: Path) -> None:
    attachment1 = _write_attachment1(tmp_path / "附件1.xlsx")
    days = (dt.date(2025, 1, 1), dt.date(2025, 1, 2))
    attachment2 = _write_wide_workbook(
        tmp_path / "附件2.xlsx",
        sheets=("实际负荷", "实际光伏"),
        days=days,
    )

    bundle = load_q2_inputs(
        attachment1_path=attachment1,
        attachment2_path=attachment2,
        load_sheet_name="实际负荷",
        pv_sheet_name="实际光伏",
    )

    assert len(bundle.fixed_prices) == 144
    assert [point.slot for point in bundle.fixed_prices] == list(range(144))
    assert len(bundle.actuals) == 288
    assert bundle.actuals[0].start == dt.datetime(2025, 1, 1, 0, 0)
    assert bundle.actuals[0].end == dt.datetime(2025, 1, 1, 0, 10)
    assert bundle.actuals[-1].slot == 143
    assert bundle.actuals[-1].end == dt.datetime(2025, 1, 3, 0, 0)
    assert bundle.fixed_prices[0].source_ref.startswith("附件1.xlsx!Sheet1!B2 ")
    assert bundle.actuals[0].load_source_ref.startswith("附件2.xlsx!实际负荷!B2 ")
    assert bundle.input_hashes == tuple(
        sorted(
            (
                ("attachment1", sha256_file(attachment1)),
                ("attachment2", sha256_file(attachment2)),
            )
        )
    )


def test_load_q2_inputs_does_not_modify_sources(tmp_path: Path) -> None:
    attachment1 = _write_attachment1(tmp_path / "附件1.xlsx")
    attachment2 = _write_wide_workbook(
        tmp_path / "附件2.xlsx",
        sheets=("实际负荷", "实际光伏"),
        days=(dt.date(2025, 1, 1),),
    )
    before = (sha256_file(attachment1), sha256_file(attachment2))

    load_q2_inputs(
        attachment1_path=attachment1,
        attachment2_path=attachment2,
        load_sheet_name="实际负荷",
        pv_sheet_name="实际光伏",
    )

    assert (sha256_file(attachment1), sha256_file(attachment2)) == before


def test_load_q2_inputs_rejects_mismatched_actual_grids(tmp_path: Path) -> None:
    day = dt.date(2025, 1, 1)
    attachment1 = _write_attachment1(tmp_path / "附件1.xlsx")
    attachment2 = _write_wide_workbook(
        tmp_path / "附件2.xlsx",
        sheets=("实际负荷", "实际光伏"),
        days=(day,),
        omitted={"实际光伏": {(day, 7)}},
    )

    with pytest.raises(InputError, match="missing_slots=.*7"):
        load_q2_inputs(
            attachment1_path=attachment1,
            attachment2_path=attachment2,
            load_sheet_name="实际负荷",
            pv_sheet_name="实际光伏",
        )


def test_load_q2_inputs_rejects_duplicate_right_endpoints(tmp_path: Path) -> None:
    attachment1 = _write_attachment1(tmp_path / "附件1.xlsx")
    attachment2 = _write_wide_workbook(
        tmp_path / "附件2.xlsx",
        sheets=("实际负荷", "实际光伏"),
        days=(dt.date(2025, 1, 1),),
        duplicate_header_slot=1,
    )

    with pytest.raises(InputError, match="duplicate key"):
        load_q2_inputs(
            attachment1_path=attachment1,
            attachment2_path=attachment2,
            load_sheet_name="实际负荷",
            pv_sheet_name="实际光伏",
        )


def test_load_q2_inputs_rejects_negative_cell_with_source_context(tmp_path: Path) -> None:
    day = dt.date(2025, 1, 1)
    attachment1 = _write_attachment1(tmp_path / "附件1.xlsx")
    attachment2 = _write_wide_workbook(
        tmp_path / "附件2.xlsx",
        sheets=("实际负荷", "实际光伏"),
        days=(day,),
        bad_values={("实际负荷", day, 4): -1.0},
    )

    with pytest.raises(InputError) as excinfo:
        load_q2_inputs(
            attachment1_path=attachment1,
            attachment2_path=attachment2,
            load_sheet_name="实际负荷",
            pv_sheet_name="实际光伏",
        )
    message = str(excinfo.value)
    assert "attachment2 load" in message
    assert "附件2.xlsx!实际负荷!F2" in message


def test_load_q2_inputs_wraps_missing_path_and_sheet(tmp_path: Path) -> None:
    attachment1 = _write_attachment1(tmp_path / "附件1.xlsx")
    attachment2 = _write_wide_workbook(
        tmp_path / "附件2.xlsx",
        sheets=("实际负荷", "实际光伏"),
        days=(dt.date(2025, 1, 1),),
    )

    with pytest.raises(InputError, match="attachment1 path is not a file"):
        load_q2_inputs(
            attachment1_path=tmp_path / "missing.xlsx",
            attachment2_path=attachment2,
            load_sheet_name="实际负荷",
            pv_sheet_name="实际光伏",
        )
    with pytest.raises(InputError, match="attachment2 load.*missing-sheet"):
        load_q2_inputs(
            attachment1_path=attachment1,
            attachment2_path=attachment2,
            load_sheet_name="missing-sheet",
            pv_sheet_name="实际光伏",
        )


def test_load_q4_2_prices_and_require_matching_grid(tmp_path: Path) -> None:
    day = dt.date(2025, 1, 1)
    attachment1 = _write_attachment1(tmp_path / "附件1.xlsx")
    attachment2 = _write_wide_workbook(
        tmp_path / "附件2.xlsx",
        sheets=("实际负荷", "实际光伏"),
        days=(day,),
    )
    attachment4 = _write_wide_workbook(
        tmp_path / "附件4.xlsx",
        sheets=("波动电价",),
        days=(day,),
    )
    q2_inputs = load_q2_inputs(
        attachment1_path=attachment1,
        attachment2_path=attachment2,
        load_sheet_name="实际负荷",
        pv_sheet_name="实际光伏",
    )

    prices = load_q4_2_prices(
        attachment4_path=attachment4,
        price_sheet_name="波动电价",
    )

    assert len(prices.prices) == 144
    assert prices.prices[0].start == dt.datetime(2025, 1, 1, 0, 0)
    assert prices.prices[-1].end == dt.datetime(2025, 1, 2, 0, 0)
    assert prices.prices[0].source_ref.startswith("附件4.xlsx!波动电价!B2 ")
    assert prices.input_hashes == (("attachment4", sha256_file(attachment4)),)
    require_matching_q4_2_grid(q2_inputs, prices)


def test_require_matching_q4_2_grid_rejects_different_days(tmp_path: Path) -> None:
    attachment1 = _write_attachment1(tmp_path / "附件1.xlsx")
    attachment2 = _write_wide_workbook(
        tmp_path / "附件2.xlsx",
        sheets=("实际负荷", "实际光伏"),
        days=(dt.date(2025, 1, 1),),
    )
    attachment4 = _write_wide_workbook(
        tmp_path / "附件4.xlsx",
        sheets=("波动电价",),
        days=(dt.date(2025, 1, 2),),
    )
    q2_inputs = load_q2_inputs(
        attachment1_path=attachment1,
        attachment2_path=attachment2,
        load_sheet_name="实际负荷",
        pv_sheet_name="实际光伏",
    )
    prices = load_q4_2_prices(
        attachment4_path=attachment4,
        price_sheet_name="波动电价",
    )

    with pytest.raises(InputError, match="Q2/Attachment 4 grid mismatch"):
        require_matching_q4_2_grid(q2_inputs, prices)


def test_historical_info_items_are_causal_at_interval_end(tmp_path: Path) -> None:
    day = dt.date(2025, 1, 1)
    attachment1 = _write_attachment1(tmp_path / "附件1.xlsx")
    attachment2 = _write_wide_workbook(
        tmp_path / "附件2.xlsx",
        sheets=("实际负荷", "实际光伏"),
        days=(day,),
    )
    attachment4 = _write_wide_workbook(
        tmp_path / "附件4.xlsx",
        sheets=("波动电价",),
        days=(day,),
    )
    q2_inputs = load_q2_inputs(
        attachment1_path=attachment1,
        attachment2_path=attachment2,
        load_sheet_name="实际负荷",
        pv_sheet_name="实际光伏",
    )
    prices = load_q4_2_prices(
        attachment4_path=attachment4,
        price_sheet_name="波动电价",
    )
    items = historical_info_items(q2_inputs, variable_prices=prices)

    before_end = info_set_at(dt.datetime(2025, 1, 1, 0, 9, 59), items)
    at_end = info_set_at(dt.datetime(2025, 1, 1, 0, 10), items)

    assert before_end.visible_items == ()
    assert [item.kind for item in at_end.visible_items] == [
        "load_actual_kw",
        "price_actual_cny_per_kwh",
        "pv_actual_kw",
    ]
    assert all(item.available_at == dt.datetime(2025, 1, 1, 0, 10) for item in at_end.visible_items)
    assert all(item.valid_time == item.available_at for item in at_end.visible_items)
    assert not any(
        item.valid_time == dt.datetime(2025, 1, 1, 0, 20) for item in at_end.visible_items
    )
