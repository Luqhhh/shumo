from __future__ import annotations

import datetime as dt
from dataclasses import replace

import pytest
from openpyxl import Workbook, load_workbook

from microgrid.artifacts import file_digest
from microgrid.problem.contracts import (
    BatteryAction,
    BatteryState,
    CaseResult,
    IntervalResult,
    apply_battery_action,
)
from microgrid.problem.q3 import REQUIRED_DECISIONS
from microgrid.problem.q3_export import (
    MAPPING_VERSION,
    create_q3_delivery,
    fill_q3_workbook,
    readback_q3_workbook,
    require_q3_export_approval,
    validate_annual_result,
)
from microgrid.problem.q3_reporting import emergency_segments, settlement_summary
from microgrid.schemas import InputError, PendingDecisionError


def _decision(repo, *, scope=True, status="approved", confirmation="test-only"):
    (repo / "configs").mkdir(exist_ok=True)
    lines = []
    for key in (*REQUIRED_DECISIONS, "D_TIME_TEMPLATE_EXPORT"):
        lines.extend(
            (
                f"[decisions.{key}]",
                f'status="{status}"',
                f'confirmed_by="{confirmation}"',
                'confirmed_at="test-only"',
                'choice="test-only"',
            )
        )
        if key == "D_TIME_TEMPLATE_EXPORT" and scope:
            lines.extend(
                (
                    'q3_scope_cases=["q3"]',
                    f'q3_mapping_version="{MAPPING_VERSION}"',
                    'q3_choice="test-only"',
                    'q3_approval_source="test-only"',
                    'q3_formal_source_run_id="source"',
                    'q3_formal_run_id="delivery"',
                    'q3_audit_sha256="wrong-test-only-hash"',
                )
            )
    (repo / "configs/decisions.toml").write_text("\n".join(lines), encoding="utf-8")


def _result(days):
    intervals = []
    state = BatteryState(6000)
    for day in days:
        for slot in range(144):
            action = BatteryAction(
                charge_kwh=1 if slot < 24 else 0, discharge_kwh=0.81 if 24 <= slot < 48 else 0
            )
            end = apply_battery_action(state, action)
            emergency = 5.0 if slot in (0, 1, 143) else 0.0
            final = 110.0 + slot
            intervals.append(
                IntervalResult(
                    day=day,
                    slot=slot,
                    load_kw=6 * (final + emergency + action.discharge_kwh - action.charge_kwh),
                    pv_kw=0,
                    planned_purchase_kwh=100,
                    adjusted_purchase_kwh=final,
                    emergency_purchase_kwh=emergency,
                    action=action,
                    state_start=state,
                    state_end=end,
                    source_ref="synthetic-test-only",
                )
            )
            state = end
    return CaseResult(
        case_id="q3",
        run_id="test-only",
        status="success",
        is_synthetic=True,
        intervals=tuple(intervals),
    )


def _workbook(days):
    wb = Workbook()
    wb.active.title = "计划购电量"
    wb.create_sheet("调整购电量")
    for name in ("计划购电量", "调整购电量"):
        for col in range(1, 148):
            wb[name].cell(1, col, f"original-header-{col}")
        for index, day in enumerate(days):
            wb[name].cell(index + 2, 1, dt.datetime.combine(day, dt.time()))
    for name, columns in (("充放电量", 6), ("紧急购电量", 3)):
        ws = wb.create_sheet(name)
        for col in range(1, columns + 1):
            ws.cell(1, col, f"header-{col}")
        ws.cell(8, 1, "⁝")
    return wb


@pytest.mark.parametrize(
    "scope,status,confirmation",
    [
        (False, "approved", "test"),
        (True, "pending", "test"),
        (True, "proposed", "test"),
        (True, "approved", ""),
    ],
)
def test_q3_export_requires_its_own_scope_and_full_confirmation(
    tmp_path, scope, status, confirmation
):
    _decision(tmp_path, scope=scope, status=status, confirmation=confirmation)
    with pytest.raises(PendingDecisionError):
        require_q3_export_approval(tmp_path)


def test_q3_supplement_gate_accepts_exact_approved_scope(tmp_path):
    _decision(tmp_path)
    assert require_q3_export_approval(tmp_path)["q3_mapping_version"] == MAPPING_VERSION


def test_q3_mapping_readback_and_original_preservation(tmp_path):
    days = (dt.date(2025, 2, 1), dt.date(2025, 2, 2))
    result = _result(days)
    wb = _workbook(days)
    template, output = tmp_path / "template.xlsx", tmp_path / "output.xlsx"
    wb.save(template)
    original_hash = file_digest(template)
    daily = {d: {"planned_cost_cny": 123.0, "adjustment_cost_cny": 0.0000002} for d in days}
    mapping = fill_q3_workbook(wb, result, daily, days)
    wb.save(output)
    wb.close()
    readback = readback_q3_workbook(template, output, result, daily, days)
    assert readback["ok"] and readback["purchase_intervals"] == 288
    assert file_digest(template) == original_hash
    assert mapping[0]["interval_start"] == "2025-02-01T00:00:00"
    assert mapping[143]["interval_end"] == "2025-02-02T00:00:00"
    assert mapping[0]["Gfinal_cell"] == "调整购电量!B2"
    assert mapping[143]["Gfinal_cell"] == "调整购电量!EO2"
    actual = load_workbook(output)
    assert actual["调整购电量"]["B2"].value == 110  # final total, not delta=10
    assert actual["调整购电量"]["EQ2"].value == pytest.approx(0.0000002)
    assert actual["充放电量"]["C2"].value == 24
    assert actual["充放电量"]["D3"].value == pytest.approx(24 * 0.81)
    assert actual["紧急购电量"]["B2"].value == "0:00-0:20"
    assert actual["紧急购电量"]["B3"].value == "23:50-0:00+1"
    actual["调整购电量"]["B2"] = 10
    actual.save(output)
    actual.close()
    with pytest.raises(InputError, match="readback mismatch"):
        readback_q3_workbook(template, output, result, daily, days)


def test_emergency_groups_respect_tolerance_and_midnight():
    result = _result((dt.date(2025, 12, 31),))
    groups = emergency_segments(result.intervals)
    assert groups[0]["energy_kwh"] == 10
    assert groups[-1]["label"] == "23:50-0:00+1"
    cleared = tuple(replace(i, emergency_purchase_kwh=1e-7) for i in result.intervals)
    assert emergency_segments(cleared) == ()
    with pytest.raises(InputError):
        emergency_segments(result.intervals[:-1])


def test_q3_settlement_report_retains_micro_net_zero_roundtrip():
    rows = tuple(
        {
            "record_type": "adjustment",
            "issue_time": f"2025-02-01T{hour}:00:00",
            "target_slot_start": "2025-02-01T18:00:00",
            "price_cny_per_kwh": 1.0,
            "cost_cny": cost,
            "delta_plus_kwh": plus,
            "delta_minus_kwh": minus,
        }
        for hour, plus, minus, cost in (("06", 1e-7, 0, 1.5e-7), ("12", 0, 1e-7, 0.5e-7))
    )
    daily, releases = settlement_summary(rows)
    assert daily[dt.date(2025, 2, 1)]["adjustment_cost_cny"] == pytest.approx(2e-7)
    assert sum(v["nonzero_events"] for v in releases.values()) == 2
    assert sum(v["material_events"] for v in releases.values()) == 0
    with pytest.raises(InputError, match="duplicate"):
        settlement_summary(rows + rows[:1])


def test_annual_export_rejects_synthetic_and_incomplete():
    result = _result((dt.date(2025, 2, 1),))
    with pytest.raises(InputError, match="nonsynthetic"):
        validate_annual_result(result)
    with pytest.raises(InputError, match="334"):
        validate_annual_result(replace(result, is_synthetic=False))


def test_q3_delivery_rejects_bad_audit_before_writing(tmp_path):
    _decision(tmp_path)
    audit = tmp_path / "audit.json"
    audit.write_text("{}", encoding="utf-8")
    with pytest.raises(InputError, match="receipt"):
        create_q3_delivery(tmp_path, source_run_id="source", run_id="delivery", audit_report=audit)
    assert not (tmp_path / "outputs/runs/q3/delivery").exists()


def test_q3_delivery_never_overwrites_existing_directory(tmp_path):
    _decision(tmp_path)
    target = tmp_path / "outputs/runs/q3/delivery"
    target.mkdir(parents=True)
    sentinel = target / "original.txt"
    sentinel.write_text("preserve", encoding="utf-8")
    with pytest.raises(InputError, match="already exists"):
        create_q3_delivery(
            tmp_path,
            source_run_id="source",
            run_id="delivery",
            audit_report=tmp_path / "missing.json",
        )
    assert sentinel.read_text() == "preserve"
