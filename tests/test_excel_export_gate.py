from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from microgrid.excel_export import TEMPLATE_EXPORT_DECISION_ID, export_case_result
from microgrid.problem.contracts import (
    BatteryAction,
    BatteryState,
    CaseResult,
    IntervalResult,
    apply_battery_action,
)
from microgrid.schemas import PendingDecisionError


def _write_export_decision(
    repo: Path, *, status: str, confirmed_by: str = "tester", confirmed_at: str = "2026-09-10"
) -> None:
    (repo / "configs").mkdir(parents=True, exist_ok=True)
    (repo / "configs" / "decisions.toml").write_text(
        "\n".join(
            [
                "[decisions.D_TIME_TEMPLATE_EXPORT]",
                f'status = "{status}"',
                'choice = ""',
                'rationale = ""',
                f'confirmed_by = "{confirmed_by}"',
                f'confirmed_at = "{confirmed_at}"',
                'source = "test"',
                "",
            ]
        ),
        encoding="utf-8",
    )


def test_template_export_is_gated_separately_from_internal_time(tmp_path):
    repo = tmp_path / "repo"
    _write_export_decision(repo, status="pending")
    result = CaseResult(case_id="q1", run_id="run-1", status="success")
    with pytest.raises(PendingDecisionError) as excinfo:
        export_case_result(repo, result)
    assert TEMPLATE_EXPORT_DECISION_ID in excinfo.value.decision_ids


def test_approved_status_without_confirmation_still_blocks_export(tmp_path):
    repo = tmp_path / "repo"
    _write_export_decision(repo, status="approved", confirmed_by="", confirmed_at="")
    result = CaseResult(case_id="q1", run_id="run-1", status="success")
    with pytest.raises(PendingDecisionError):
        export_case_result(repo, result)


def _q1_case_result() -> CaseResult:
    day = dt.date(2025, 1, 1)
    state = BatteryState(6000.0)
    intervals: list[IntervalResult] = []
    for slot in range(144):
        start = state
        action = BatteryAction()
        end = apply_battery_action(start, action)
        intervals.append(
            IntervalResult(
                day=day,
                slot=slot,
                load_kw=600.0,
                pv_kw=0.0,
                planned_purchase_kwh=100.0,
                adjusted_purchase_kwh=100.0,
                emergency_purchase_kwh=0.0,
                action=action,
                state_start=start,
                state_end=end,
            )
        )
        state = end
    return CaseResult(
        case_id="q1",
        run_id="run-1",
        status="success",
        intervals=tuple(intervals),
    )


def test_approved_template_writer_exports_q1_and_reads_back(synthetic_template_repo):
    repo = synthetic_template_repo
    _write_export_decision(repo, status="approved")
    result = _q1_case_result()
    output = export_case_result(repo, result, output_path=repo / "out_result1.xlsx")
    assert output.is_file()
    wb = load_workbook(output, data_only=False)
    try:
        plan = wb["计划购电量"]
        charge = wb["充放电量"]
        assert plan["A2"].value == "0:10-0:20"
        assert plan["A145"].value == "0:00+1-0:10+1"
        assert plan["B2"].value == pytest.approx(100.0)
        assert charge["E2"].value == pytest.approx(6000.0)
        assert charge["E3"].value == pytest.approx(6000.0)
    finally:
        wb.close()


def _q2_case_result() -> CaseResult:
    state = BatteryState(6000.0)
    intervals: list[IntervalResult] = []
    for day_index in range(2):
        day = dt.date(2025, 2, 1) + dt.timedelta(days=day_index)
        for slot in range(144):
            start = state
            if slot == 0:
                action = BatteryAction(charge_kwh=10.0)
            elif slot == 1:
                action = BatteryAction(discharge_kwh=8.1)
            else:
                action = BatteryAction()
            end = apply_battery_action(start, action)
            intervals.append(
                IntervalResult(
                    day=day,
                    slot=slot,
                    load_kw=600.0,
                    pv_kw=0.0,
                    planned_purchase_kwh=float(slot + 1 + day_index),
                    adjusted_purchase_kwh=float(slot + 1 + day_index),
                    emergency_purchase_kwh=(
                        3.0
                        if (day_index, slot) == (0, 0)
                        else 5.0
                        if (day_index, slot) == (1, 143)
                        else 0.0
                    ),
                    action=action,
                    state_start=start,
                    state_end=end,
                )
            )
            state = end
    return CaseResult(
        case_id="q2",
        run_id="run-2",
        status="success",
        intervals=tuple(intervals),
    )


def test_approved_template_writer_exports_q2_and_reads_back(synthetic_template_repo):
    repo = synthetic_template_repo
    _write_export_decision(repo, status="approved")
    template = repo / "data" / "templates" / "result2.xlsx"
    wb = load_workbook(template)
    try:
        plan = wb["计划购电量"]
        plan.cell(row=2, column=1, value=dt.datetime(2025, 2, 1))
        plan.cell(row=3, column=1, value=dt.datetime(2025, 2, 2))
        wb.save(template)
    finally:
        wb.close()
    snapshot = repo / "outputs" / "runs" / "q2" / "run-2" / "input_snapshot.json"
    snapshot.parent.mkdir(parents=True)
    snapshot.write_text(
        json.dumps(
            {
                "fixed_prices": [
                    {"slot": slot, "price_cny_per_kwh": 0.4 + slot / 1000.0} for slot in range(144)
                ]
            }
        ),
        encoding="utf-8",
    )

    result = _q2_case_result()
    output = export_case_result(repo, result, output_path=repo / "out_result2.xlsx")

    wb = load_workbook(output, data_only=False)
    try:
        plan = wb["计划购电量"]
        assert plan["B2"].value == pytest.approx(1.0)
        assert plan["EO2"].value == pytest.approx(144.0)
        assert plan["EP2"].value == pytest.approx(sum(range(1, 145)))
        assert plan["EQ2"].value == pytest.approx(
            sum((slot + 1) * (0.4 + slot / 1000.0) for slot in range(144))
        )

        charge = wb["充放电量"]
        assert charge.max_row == 13
        assert charge["A2"].value.date() == dt.date(2025, 2, 1)
        assert charge["A8"].value.date() == dt.date(2025, 2, 2)
        assert charge["C2"].value == pytest.approx(10.0)
        assert charge["D2"].value == pytest.approx(8.1)
        assert charge["E2"].value == dt.time(0, 0)
        assert charge["E3"].value == "24:00"
        assert charge["F2"].value == pytest.approx(6000.0)
        assert charge["F3"].value == pytest.approx(6000.0)

        emergency = wb["紧急购电量"]
        assert emergency.max_row == 3
        assert emergency["A2"].value.date() == dt.date(2025, 2, 1)
        assert emergency["B2"].value == "0:10-0:20"
        assert emergency["C2"].value == pytest.approx(3.0)
        assert emergency["A3"].value.date() == dt.date(2025, 2, 2)
        assert emergency["B3"].value == "0:00-0:10+1"
        assert emergency["C3"].value == pytest.approx(5.0)
    finally:
        wb.close()


def test_q2_template_writer_is_the_default_for_q2(synthetic_template_repo):
    repo = synthetic_template_repo
    _write_export_decision(repo, status="approved")
    template = repo / "data" / "templates" / "result2.xlsx"
    wb = load_workbook(template)
    try:
        plan = wb["计划购电量"]
        plan.cell(row=2, column=1, value=dt.datetime(2025, 2, 1))
        plan.cell(row=3, column=1, value=dt.datetime(2025, 2, 2))
        wb.save(template)
    finally:
        wb.close()
    snapshot = repo / "outputs" / "runs" / "q2" / "run-2" / "input_snapshot.json"
    snapshot.parent.mkdir(parents=True)
    snapshot.write_text(
        json.dumps(
            {"fixed_prices": [{"slot": slot, "price_cny_per_kwh": 1.0} for slot in range(144)]}
        ),
        encoding="utf-8",
    )

    output = export_case_result(repo, _q2_case_result())

    assert output.name == "result2.xlsx"
    assert output.parent == repo / "outputs" / "runs" / "q2" / "run-2" / "results"
