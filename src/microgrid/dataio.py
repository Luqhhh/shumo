"""Safe ingestion, raw-file inventory and lossless readers.

The module deliberately preserves raw labels and cell values.  It never fills
power/price/forecast values and never decides what a time label means for energy.
"""

from __future__ import annotations

import datetime as _dt
import hashlib
import json
import os
import re
import shutil
import stat
import zipfile
from collections import Counter
from collections.abc import Iterator
from pathlib import Path, PurePosixPath
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]

from .schemas import (
    ForecastRecord,
    InputError,
    ObservationRecord,
    TimeLabelError,
    VersionConflictError,
)
from .timekeys import (
    issue_time_from_forecast_label,
    parse_interval,
    parse_time_label,
    valid_time,
)

MAX_ZIP_UNCOMPRESSED = 200 * 1024 * 1024
MAX_ZIP_RATIO = 250.0
FORECAST_LEAD_RE = re.compile(r"预报\s*(?P<n>\d{1,2})\s*小时")


def utc_now() -> str:
    return _dt.datetime.now(_dt.UTC).isoformat(timespec="seconds")


def sha256_file(path: str | Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as fh:
        while True:
            chunk = fh.read(chunk_size)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def ensure_dir(path: str | Path) -> Path:
    p = Path(path)
    p.mkdir(parents=True, exist_ok=True)
    return p


def _norm_rel(path: str) -> PurePosixPath:
    return PurePosixPath(str(path).replace("\\", "/"))


def _is_safe_rel(rel: PurePosixPath) -> bool:
    if rel.is_absolute():
        return False
    if not rel.parts:
        return False
    if any(part in {"", ".", ".."} for part in rel.parts):
        return False
    if rel.parts[0].endswith(":"):
        return False
    return True


def decode_zip_name(info: zipfile.ZipInfo) -> str:
    """Decode Chinese ZIP names without destructive guessing.

    The caller still treats the decoded name as a candidate; ambiguous names
    are reported, not silently assumed to be a particular official file.
    """

    if info.flag_bits & 0x800:
        return info.filename
    try:
        raw = info.filename.encode("cp437")
    except UnicodeEncodeError:
        return info.filename
    for encoding in ("utf-8", "gb18030"):
        try:
            candidate = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        if "\ufffd" not in candidate:
            return candidate
    return info.filename


def classify_input(rel: PurePosixPath) -> tuple[str, str] | None:
    """Map an official C-problem relative path to (kind, repository path).

    Returns None for unrelated A/B/D/E files and incidental OS metadata.
    """

    name = rel.name
    lower = name.lower()
    parts = [p.lower() for p in rel.parts]

    if name.endswith(":Zone.Identifier") or lower.endswith(".zone.identifier"):
        return None
    if name.startswith("~$") or name.startswith("."):
        return None

    if name == "C题.pdf" or lower == "c题.pdf" or name == "C题.PDF":
        return "problem", "resources/C题.pdf"

    if lower == "format2026.doc":
        return "rules", "resources/format2026.doc"
    if "咨询问题解答" in name or "人工智能工具使用规定" in name or "数模竞赛注意事项" in name:
        return "rules", f"resources/{name}"

    if lower.endswith(".xlsx"):
        if re.fullmatch(r"result4-[23]\.xlsx", lower) or re.fullmatch(r"result[123]\.xlsx", lower):
            return "template", f"data/templates/{name}"
        if any("附件5" in p for p in parts):
            return None
        if re.fullmatch(r"附件[1-4](?:\.xlsx)?", name) or re.fullmatch(r"附件[1-4]\.xlsx", lower):
            return "raw", f"data/raw/{name}"
        if name in {"附件1.xlsx", "附件2.xlsx", "附件3.xlsx", "附件4.xlsx"}:
            return "raw", f"data/raw/{name}"

    return None


def _zip_infos(zf: zipfile.ZipFile) -> list[tuple[zipfile.ZipInfo, str, PurePosixPath]]:
    out: list[tuple[zipfile.ZipInfo, str, PurePosixPath]] = []
    total_uncompressed = 0
    for info in zf.infolist():
        name = decode_zip_name(info)
        rel = _norm_rel(name)
        if info.is_dir():
            continue
        if not _is_safe_rel(rel):
            raise InputError(f"unsafe ZIP member path: {name!r}")
        mode = (info.external_attr >> 16) & 0xFFFF
        if stat.S_ISLNK(mode):
            raise InputError(f"ZIP symlink is not allowed: {name!r}")
        total_uncompressed += info.file_size
        if total_uncompressed > MAX_ZIP_UNCOMPRESSED:
            raise InputError("ZIP uncompressed size exceeds safety limit")
        if info.file_size > 0 and info.compress_size > 0:
            if info.file_size / info.compress_size > MAX_ZIP_RATIO:
                raise InputError(f"ZIP compression ratio is suspiciously high: {name!r}")
        out.append((info, name, rel))
    return out


def _iter_source_files(
    source: Path, *, exclude_root: Path | None = None
) -> Iterator[tuple[str, Path]]:
    if not source.exists():
        raise InputError(f"source does not exist: {source}")
    if source.is_file() and source.suffix.lower() == ".zip":
        with zipfile.ZipFile(source) as zf:
            for _info, name, _rel in _zip_infos(zf):
                yield name, source  # caller opens zip lazily; handled below
        return
    if not source.is_dir():
        raise InputError(f"source must be a directory or .zip file: {source}")
    for path in sorted(source.rglob("*")):
        if not path.is_file():
            continue
        if any(part in {".git", "__pycache__", ".venv"} for part in path.parts):
            continue
        if exclude_root is not None:
            try:
                path.resolve().relative_to(exclude_root)
            except ValueError:
                pass
            else:
                continue
        yield str(path.relative_to(source)), path


def _open_zip_member(source: Path, name: str) -> bytes:
    with zipfile.ZipFile(source) as zf:
        for info in zf.infolist():
            if decode_zip_name(info) == name:
                return zf.read(info)
    raise InputError(f"ZIP member disappeared while reading: {name}")


def _write_bytes_atomic(target: Path, data: bytes) -> None:
    ensure_dir(target.parent)
    tmp = target.with_name(target.name + ".tmp-ingest")
    tmp.write_bytes(data)
    os.replace(tmp, target)


def ingest_source(source_path: str | Path, repo_root: str | Path) -> dict[str, Any]:
    """Copy only C-problem files into the Stage 0 repository layout.

    The source is never modified.  Same path + same hash is idempotent; same
    path + different hash raises `VersionConflictError` after writing a report.
    """

    source = Path(source_path)
    repo = Path(repo_root)
    ensure_dir(repo)
    items: list[dict[str, Any]] = []
    conflicts: list[dict[str, Any]] = []
    is_zip = source.is_file() and source.suffix.lower() == ".zip"

    if is_zip:
        with zipfile.ZipFile(source) as zf:
            entries = list(_zip_infos(zf))
        for _info, name, rel in entries:
            classified = classify_input(rel)
            if classified is None:
                continue
            kind, dest_rel = classified
            data = _open_zip_member(source, name)
            digest = hashlib.sha256(data).hexdigest()
            target = repo / dest_rel
            entry = _record_import(kind, name, dest_rel, len(data), digest, target)
            if entry["status"] == "conflict":
                conflicts.append(entry)
            else:
                _write_bytes_atomic(target, data)
                if sha256_file(target) != digest:
                    raise InputError(f"copy verification failed for {dest_rel}")
            items.append(entry)
        source_label = source.name
    else:
        for rel_str, path in _iter_source_files(source, exclude_root=repo.resolve()):
            rel = _norm_rel(rel_str)
            classified = classify_input(rel)
            if classified is None:
                continue
            kind, dest_rel = classified
            digest_before = sha256_file(path)
            size = path.stat().st_size
            target = repo / dest_rel
            entry = _record_import(kind, rel_str, dest_rel, size, digest_before, target)
            if entry["status"] == "conflict":
                conflicts.append(entry)
            else:
                ensure_dir(target.parent)
                shutil.copyfile(path, target)
                if sha256_file(target) != digest_before:
                    raise InputError(f"copy verification failed for {dest_rel}")
                if sha256_file(path) != digest_before:
                    raise InputError(f"source file changed while reading: {rel_str}")
            items.append(entry)
        source_label = source.name

    manifest = {
        "schema_version": 1,
        "generated_at": utc_now(),
        "source": {
            "kind": "zip" if is_zip else "directory",
            "label": source_label,
            "file_count_considered": len(items),
        },
        "items": items,
        "conflicts": conflicts,
        "note": (
            "Only C-problem and official rule files are imported. "
            "Absolute source paths are intentionally not stored."
        ),
    }
    manifest_path = repo / "records" / "inputs_manifest.json"
    ensure_dir(manifest_path.parent)
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    if conflicts:
        joined = "; ".join(f"{c['repository_path']} ({c['sha256'][:12]})" for c in conflicts)
        raise VersionConflictError(f"same path with different hash, no silent overwrite: {joined}")
    return manifest


def _record_import(
    kind: str,
    source_name: str,
    repository_path: str,
    size: int,
    digest: str,
    target: Path,
) -> dict[str, Any]:
    status = "imported"
    existing_hash: str | None = None
    if target.exists():
        existing_hash = sha256_file(target)
        status = "unchanged" if existing_hash == digest else "conflict"
    return {
        "kind": kind,
        "source_name": source_name,
        "repository_path": repository_path,
        "size": size,
        "sha256": digest,
        "existing_sha256": existing_hash,
        "imported_at": utc_now(),
        "status": status,
    }


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _has_time_label_syntax(value: str) -> bool:
    if "-" in value and ":" in value:
        return True
    if "+" in value and ":" in value:
        return True
    return bool(re.fullmatch(r"\d{1,2}:\d{2}", value.strip()))


def inspect_workbook(path: Path) -> dict[str, Any]:
    """Read-only structural inventory of one workbook."""

    wb = load_workbook(path, data_only=False, read_only=False)
    try:
        result: dict[str, Any] = {
            "file": str(path),
            "sheet_names": list(wb.sheetnames),
            "epoch": wb.epoch.isoformat() if getattr(wb, "epoch", None) else None,
            "sheets": [],
        }
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            max_row = ws.max_row or 0
            max_col = ws.max_column or 0
            non_empty = 0
            formulas = 0
            type_counts: Counter[str] = Counter()
            first_non_empty_row: int | None = None
            header_values: list[str] = []
            first_col_values: list[str] = []
            duplicate_first_col: list[str] = []
            first_col_seen: set[str] = set()
            time_label_candidates: list[str] = []
            for row_idx, row in enumerate(ws.iter_rows(), start=1):
                row_non_empty = 0
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    if isinstance(value, str) and value == "":
                        non_empty += 1
                        type_counts["empty_string"] += 1
                        continue
                    non_empty += 1
                    row_non_empty += 1
                    if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                        formulas += 1
                    if isinstance(value, bool):
                        type_counts["bool"] += 1
                    elif isinstance(value, (int, float)):
                        type_counts["number"] += 1
                    elif isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
                        type_counts[type(value).__name__] += 1
                    else:
                        type_counts["string"] += 1
                    if cell.column == 1 and row_idx > 1 and not isinstance(value, str):
                        text = _cell_text(value)
                        if text and text not in first_col_seen:
                            first_col_seen.add(text)
                        elif text in first_col_seen and text not in duplicate_first_col:
                            duplicate_first_col.append(text)
                    if isinstance(value, str) and _has_time_label_syntax(value):
                        if len(time_label_candidates) < 1000:
                            time_label_candidates.append(value)
                if row_non_empty and first_non_empty_row is None:
                    first_non_empty_row = row_idx
                    header_values = [_cell_text(c.value) for c in row[: min(max_col, 20)]]
            for row in ws.iter_rows(min_col=1, max_col=1, max_row=min(max_row, 8)):
                for cell in row:
                    if cell.value is not None:
                        first_col_values.append(_cell_text(cell.value))
            merged = [str(rng) for rng in ws.merged_cells.ranges]
            result["sheets"].append(
                {
                    "name": sheet_name,
                    "max_row": max_row,
                    "max_column": max_col,
                    "dimension": ws.calculate_dimension(),
                    "non_empty_cells": non_empty,
                    "formula_cells": formulas,
                    "type_counts": dict(type_counts),
                    "first_non_empty_row": first_non_empty_row,
                    "header_values": header_values,
                    "first_column_values": first_col_values,
                    "duplicate_first_column_values": duplicate_first_col[:20],
                    "time_label_candidates": time_label_candidates,
                    "merged_ranges": merged,
                    "freeze_panes": ws.freeze_panes,
                }
            )
        return result
    finally:
        wb.close()


def inventory(repo_root: str | Path) -> dict[str, Any]:
    repo = Path(repo_root)
    raw_files = sorted((repo / "data" / "raw").glob("*.xlsx"))
    template_files = sorted((repo / "data" / "templates").glob("*.xlsx"))
    out: dict[str, Any] = {
        "schema_version": 1,
        "generated_at": utc_now(),
        "raw_directory": str((repo / "data" / "raw").relative_to(repo)),
        "template_directory": str((repo / "data" / "templates").relative_to(repo)),
        "raw_files": [],
        "template_files": [],
        "warnings": [],
    }
    for path in raw_files:
        rec = inspect_workbook(path)
        rec["file"] = str(path.relative_to(repo))
        out["raw_files"].append(rec)
    for path in template_files:
        rec = inspect_workbook(path)
        rec["file"] = str(path.relative_to(repo))
        # Literal interval anomalies are reported, never repaired.
        for sheet in rec["sheets"]:
            for label in sheet.get("time_label_candidates", []):
                if not isinstance(label, str) or "-" not in label:
                    continue
                try:
                    parsed = parse_interval(label)
                except TimeLabelError:
                    continue
                if parsed.suspicious:
                    out["warnings"].append(
                        {
                            "file": path.name,
                            "sheet": sheet["name"],
                            "label": label,
                            "reason": parsed.reason,
                            "span_minutes": parsed.span_minutes,
                        }
                    )
        out["template_files"].append(rec)
    return out


def write_inventory(repo_root: str | Path, inventory_data: dict[str, Any] | None = None) -> Path:
    repo = Path(repo_root)
    data = inventory_data or inventory(repo)
    out_dir = ensure_dir(repo / "outputs" / "inventory")
    json_path = out_dir / "inventory.json"
    json_path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    md_path = out_dir / "inventory.md"
    lines = [
        "# Data inventory",
        "",
        f"- generated_at: {data.get('generated_at', '')}",
        f"- raw files: {len(data.get('raw_files', []))}",
        f"- template files: {len(data.get('template_files', []))}",
        f"- warnings: {len(data.get('warnings', []))}",
        "",
    ]
    for rec in data.get("raw_files", []) + data.get("template_files", []):
        lines.append(f"## {Path(rec['file']).name}")
        lines.append(f"- sheets: {', '.join(rec['sheet_names'])}")
        for sheet in rec["sheets"]:
            lines.append(
                f"- {sheet['name']}: {sheet['dimension']}, "
                f"non_empty={sheet['non_empty_cells']}, formulas={sheet['formula_cells']}"
            )
        lines.append("")
    if data.get("warnings"):
        lines.append("## Literal time-label warnings")
        for warning in data["warnings"]:
            lines.append(
                f"- {warning['file']} / {warning['sheet']}: {warning['label']} "
                f"({warning['reason']})"
            )
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return json_path


def read_attachment1(path: str | Path) -> list[ObservationRecord]:
    """Lossless reader for 附件1; no date is invented for the clock column."""

    p = Path(path)
    digest = sha256_file(p)
    wb = load_workbook(p, data_only=False)
    try:
        ws = wb["Sheet1"]
        records: list[ObservationRecord] = []
        for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            raw_time = row[0].value
            if raw_time is None:
                continue
            parsed = parse_time_label(raw_time, base_date=None)
            for col_idx, (value, unit, kind) in enumerate(
                [
                    (row[1].value if len(row) > 1 else None, "元/kWh", "price_cny_per_kwh"),
                    (row[2].value if len(row) > 2 else None, "kW", "load_kw"),
                    (row[3].value if len(row) > 3 else None, "kW", "pv_forecast_kw"),
                ],
                start=2,
            ):
                if value is None:
                    continue
                col_letter = ws.cell(row=idx, column=col_idx).column_letter
                records.append(
                    ObservationRecord(
                        source_file=str(p),
                        sheet_name=ws.title,
                        cell_ref=f"{col_letter}{idx}",
                        source_hash=digest,
                        raw_time_label=str(raw_time),
                        source_date=None,
                        parsed_day_offset=parsed.day_offset,
                        parsed_minute_of_day=parsed.minute_of_day,
                        parsed_timestamp=parsed.as_datetime(),
                        value=value,
                        unit=unit,
                        kind=kind,
                    )
                )
        return records
    finally:
        wb.close()


def _resolve_time_label(value: Any, *, base_date: _dt.date | None = None) -> Any:
    return parse_time_label(value, base_date=base_date)


def read_wide_attachment(
    path: str | Path,
    *,
    sheet_name: str,
    kind: str,
    unit: str,
) -> list[dict[str, Any]]:
    """Read an attachment2/4-style date x time wide table into raw records."""

    p = Path(path)
    digest = sha256_file(p)
    wb = load_workbook(p, data_only=False)
    try:
        ws = wb[sheet_name]
        header_cells = [
            (cell.column, cell.value)
            for cell in ws[1]
            if cell.column > 1 and cell.value is not None
        ]
        records: list[dict[str, Any]] = []
        for row in ws.iter_rows(min_row=2):
            if not row:
                continue
            date_value = row[0].value
            if date_value is None:
                continue
            date_record = parse_time_label(date_value)
            source_date = date_record.base_date or date_record.as_datetime()
            for col_idx, header_value in header_cells:
                if col_idx > len(row):
                    continue
                value = row[col_idx - 1].value
                if value is None:
                    continue
                parsed = _resolve_time_label(header_value, base_date=date_record.base_date)
                records.append(
                    {
                        "source_file": str(p),
                        "sheet_name": sheet_name,
                        "cell_ref": f"{ws.cell(row=row[0].row, column=col_idx).column_letter}{row[0].row}",
                        "source_hash": digest,
                        "source_date": source_date.isoformat()
                        if hasattr(source_date, "isoformat")
                        else str(source_date),
                        "raw_time_label": str(header_value),
                        "parsed_day_offset": parsed.day_offset,
                        "parsed_minute_of_day": parsed.minute_of_day,
                        "parsed_timestamp": (
                            _dt.datetime.combine(
                                date_record.base_date or _dt.date(1900, 1, 1), _dt.time()
                            )
                            + _dt.timedelta(
                                days=parsed.day_offset,
                                minutes=parsed.minute_of_day,
                            )
                        ).isoformat()
                        if date_record.base_date is not None
                        else None,
                        "value": value,
                        "unit": unit,
                        "kind": kind,
                    }
                )
        return records
    finally:
        wb.close()


def read_forecasts(path: str | Path) -> list[ForecastRecord]:
    """Read 附件3 without forward-filling any forecast value.

    Only the date cell at the start of a 4-row block is propagated to that
    block when later date cells are empty strings.  Values are never filled.
    """

    p = Path(path)
    digest = sha256_file(p)
    wb = load_workbook(p, data_only=False)
    try:
        ws = wb["Sheet1"]
        records: list[ForecastRecord] = []
        current_date_value: Any = None
        for row in ws.iter_rows(min_row=2):
            date_value = row[0].value if len(row) > 0 else None
            if not (date_value is None or (isinstance(date_value, str) and not date_value.strip())):
                current_date_value = date_value
            if current_date_value is None:
                continue
            clock_value = row[1].value if len(row) > 1 else None
            if clock_value is None:
                continue
            try:
                issue = issue_time_from_forecast_label(current_date_value, clock_value)
            except TimeLabelError:
                continue
            for cell in row[2:]:
                header = ws.cell(row=1, column=cell.column).value
                match = FORECAST_LEAD_RE.search(str(header or ""))
                if not match or cell.value is None:
                    continue
                lead = int(match.group("n"))
                try:
                    valid = valid_time(issue, lead)
                except TimeLabelError:
                    continue
                records.append(
                    ForecastRecord(
                        issue_time=issue,
                        valid_time=valid,
                        lead_hours=lead,
                        pv_forecast_kw=float(cell.value),
                        source_ref=(
                            f"{p.name}!{ws.title}!{ws.cell(row=cell.row, column=cell.column).coordinate}"
                            f" sha256={digest[:12]}"
                        ),
                    )
                )
        return records
    finally:
        wb.close()


def load_inputs_manifest(repo_root: str | Path) -> dict[str, Any]:
    path = Path(repo_root) / "records" / "inputs_manifest.json"
    if not path.exists():
        return {"schema_version": 1, "items": [], "items_available": False}
    data = json.loads(path.read_text(encoding="utf-8"))
    data["items_available"] = True
    return data
