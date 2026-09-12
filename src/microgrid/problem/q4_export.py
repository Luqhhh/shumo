"""Full-year Q4 template writer and independent cell readback."""

from __future__ import annotations

import datetime as dt
import json
import shutil
from copy import copy
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..approvals import load_decisions
from ..dataio import sha256_file
from ..schemas import InputError
from ..timekeys import parse_time_label
from .contracts import BatteryAction, BatteryState, CaseResult
from .dispatch_feedback import ExecutionRecord
from .purchase_ledger import Bill, ContractVersion, PurchaseLedger
from .q4_common import (
    ACTION_START,
    BIAS_MODEL_VERSION,
    BLEND_DECISIONS,
    BLEND_MODEL_VERSION,
    STEP,
    YEAR_END,
    atomic_json,
    reserve_start_from_config,
)
from .q4_evidence import (
    audit_controller_chain,
    check_inventory,
    check_model_binding,
    supplementary_dir,
)
from .q4_inputs import load_q4_inputs
from .q4_performance import Performance
from .q4_validation import validate_q4_run
from .result_io import case_result_to_dict


def load_replay_evidence(run_dir: Path, case_id: str):
    for name in ("forecasts", "solver_records", "contracts", "cost_ledger", "execution_feedback"):
        if not (run_dir / f"{name}.jsonl").is_file():
            raise InputError(f"missing Q4 replay evidence: {name}.jsonl")
    ledger = PurchaseLedger(case_id)
    executions = []
    for name in ("contracts", "cost_ledger", "execution_feedback"):
        with (run_dir / f"{name}.jsonl").open(encoding="utf-8") as stream:
            for line in stream:
                row = json.loads(line)
                if name == "contracts":
                    row["day"] = dt.date.fromisoformat(row["day"])
                    row["event_time"] = dt.datetime.fromisoformat(row["event_time"])
                    for key in ("quantities", "plus", "minus"):
                        row[key] = tuple(row[key])
                    ledger.versions.setdefault(row["day"], []).append(ContractVersion(**row))
                elif name == "cost_ledger":
                    row["slot_start"] = dt.datetime.fromisoformat(row["slot_start"])
                    bill = Bill(**row)
                    if bill.slot_start in ledger.bills:
                        raise InputError("duplicate cost ledger record")
                    ledger.bills[bill.slot_start] = bill
                else:
                    data = row["execution"]
                    for key in ("action", "intent"):
                        data[key] = BatteryAction(**data[key])
                    for key in ("state_start", "state_end"):
                        data[key] = BatteryState(**data[key])
                    data["reasons"] = tuple(data["reasons"])
                    executions.append(ExecutionRecord(**data))
    for versions in ledger.versions.values():
        for index, version in enumerate(versions):
            if version.version:
                bill = ledger.bills.get(version.event_time)
                if bill is None:
                    raise InputError("missing actual trade bill")
                versions[index] = replace(version, trade_price=bill.execution_price)
    return ledger, tuple(executions)


def emergency_segments(rows):
    segments = []
    first = None
    total = 0.0
    for k in range(145):
        active = k < 144 and rows[k].emergency_purchase_kwh > 1e-6
        if active:
            if first is None:
                first = k
            total += rows[k].emergency_purchase_kwh
        elif first is not None:

            def endpoint(slot):
                return "0:00+1" if slot == 144 else f"{slot // 6}:{slot % 6 * 10:02d}"

            segments.append((f"{endpoint(first)}-{endpoint(k)}", total, first, k))
            first, total = None, 0.0
    return segments or [(None, 0.0, 0, 144)]


def fill_workbook(wb, result: CaseResult, ledger: PurchaseLedger):
    """Populate a COPY, returning explicit source-slot/cell mappings."""
    plan_sheet = wb["计划购电量"]
    charge_sheet, emergency_sheet = wb["充放电量"], wb["紧急购电量"]
    charge_styles = [copy(charge_sheet.cell(2, col)._style) for col in range(1, 7)]
    emergency_styles = [copy(emergency_sheet.cell(2, col)._style) for col in range(1, 4)]
    for sheet in (charge_sheet, emergency_sheet):
        for merged in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged))
        sheet.delete_rows(2, sheet.max_row)
    mappings, expected = [], []
    emergency_row = 2
    for i in range(334):
        rows = result.intervals[i * 144 : (i + 1) * 144]
        day = (ACTION_START + dt.timedelta(days=i)).date()
        if len(rows) != 144 or rows[0].day != day:
            raise InputError("Q4 export requires exactly 334 complete natural days")
        date_value = plan_sheet.cell(i + 2, 1).value
        expected_day = str(day)
        if date_value is None:
            raise InputError("Q4 template missing date")
        if parse_time_label(date_value).base_date != day:
            raise InputError("Q4 template date order differs from annual trajectory")
        # Keep the template's original date labels and interval headers.
        for k, row in enumerate(rows):
            column = k + 2
            plan_sheet.cell(i + 2, column, row.planned_purchase_kwh)
            expected.append(("计划购电量", i + 2, column, row.planned_purchase_kwh, 1e-6))
            mapping = {
                "day": expected_day,
                "slot": k,
                "interval_start": str(ACTION_START + i * dt.timedelta(days=1) + k * STEP),
                "interval_end": str(ACTION_START + i * dt.timedelta(days=1) + (k + 1) * STEP),
                "G0_cell": f"计划购电量!{get_column_letter(column)}{i + 2}",
                "charge_block_cell": f"充放电量!C{2 + 6 * i + k // 24}",
                "discharge_block_cell": f"充放电量!D{2 + 6 * i + k // 24}",
            }
            if result.case_id == "q4_3":
                grid_final = ledger.versions[day][-1].quantities[k]
                wb["调整购电量"].cell(i + 2, column, grid_final)
                expected.append(("调整购电量", i + 2, column, grid_final, 1e-6))
                mapping["Gfinal_cell"] = f"调整购电量!{get_column_letter(column)}{i + 2}"
            mappings.append(mapping)
        bills = [
            ledger.bills[ACTION_START + i * dt.timedelta(days=1) + k * STEP] for k in range(144)
        ]
        planned_cost = sum(b.planned_cost_cny for b in bills)
        adjustment_cost = sum(b.adjustment_cost_cny for b in bills)
        for sheet_name, quantity, cost in (
            ("计划购电量", sum(r.planned_purchase_kwh for r in rows), planned_cost),
            *(
                [("调整购电量", sum(ledger.versions[day][-1].quantities), adjustment_cost)]
                if result.case_id == "q4_3"
                else []
            ),
        ):
            wb[sheet_name].cell(i + 2, 146, quantity)
            wb[sheet_name].cell(i + 2, 147, cost)
            expected.extend(
                ((sheet_name, i + 2, 146, quantity, 1e-6), (sheet_name, i + 2, 147, cost, 0.01))
            )
        start_row = 2 + 6 * i
        for block in range(6):
            target_row = start_row + block
            for col in range(1, 7):
                charge_sheet.cell(target_row, col)._style = copy(charge_styles[col - 1])
            charge_sheet.cell(target_row, 2, f"{block * 4}:00-{(block + 1) * 4}:00")
            expected.append(
                ("充放电量", target_row, 2, f"{block * 4}:00-{(block + 1) * 4}:00", None)
            )
            for col, amount in (
                (3, sum(r.action.charge_kwh for r in rows[block * 24 : (block + 1) * 24])),
                (4, sum(r.action.discharge_kwh for r in rows[block * 24 : (block + 1) * 24])),
            ):
                charge_sheet.cell(target_row, col, amount)
                expected.append(("充放电量", target_row, col, amount, 1e-6))
        charge_sheet.cell(start_row, 1, date_value)
        expected.append(("充放电量", start_row, 1, date_value, None))
        charge_sheet.merge_cells(
            start_row=start_row, end_row=start_row + 5, start_column=1, end_column=1
        )
        for offset, label, energy in (
            (0, "0:00", rows[0].state_start.energy_kwh),
            (1, "24:00", rows[-1].state_end.energy_kwh),
        ):
            charge_sheet.cell(start_row + offset, 5, label)
            expected.append(("充放电量", start_row + offset, 5, label, None))
            charge_sheet.cell(start_row + offset, 6, energy)
            expected.append(("充放电量", start_row + offset, 6, energy, 1e-6))
        for label, amount, first, last in emergency_segments(rows):
            for col in range(1, 4):
                emergency_sheet.cell(emergency_row, col)._style = copy(emergency_styles[col - 1])
            emergency_sheet.cell(emergency_row, 1, date_value)
            emergency_sheet.cell(emergency_row, 2, label)
            emergency_sheet.cell(emergency_row, 3, amount)
            expected.extend(
                (
                    ("紧急购电量", emergency_row, 1, date_value, None),
                    ("紧急购电量", emergency_row, 2, label, None),
                )
            )
            expected.append(("紧急购电量", emergency_row, 3, amount, 1e-6))
            for slot_index in range(first, last):
                mappings[i * 144 + slot_index]["emergency_cell"] = f"紧急购电量!C{emergency_row}"
            emergency_row += 1
    return mappings, expected


def _export_q4(repo: Path, result: CaseResult, output: Path, performance: Performance) -> Path:
    if result.status != "success" or result.is_synthetic:
        raise InputError("Q4 formal export requires successful nonsynthetic annual main result")
    run_dir = repo / "outputs" / "runs" / result.case_id / result.run_id
    if output.exists():
        raise InputError("Q4 export output already exists")
    if not (run_dir / "evidence_manifest.json").is_file():
        raise InputError(
            "Historical Q4 run: preserve its existing export and verify supplementary evidence; a new export requires runtime plan evidence"
        )
    domain_path = run_dir / "domain_result.json"
    if not domain_path.is_file() or json.loads(domain_path.read_text()) != case_result_to_dict(
        result
    ):
        raise InputError("Q4 export result differs from saved source")
    evidence_issues = check_model_binding(repo, run_dir) + check_inventory(repo, run_dir)
    if evidence_issues:
        raise InputError("Q4 export evidence failed: " + ";".join(evidence_issues[:8]))
    ledger, executions = load_replay_evidence(run_dir, result.case_id)
    with performance.measure("input_preflight"):
        inputs = load_q4_inputs(repo, result.case_id)
    snapshot = json.loads((run_dir / "input_snapshot.json").read_text())
    if inputs.source_hashes != snapshot["source_hashes"]:
        raise InputError("Q4 source hashes differ from run inputs")
    config = json.loads((run_dir / "effective_config.json").read_text())
    plans_path = run_dir / "dispatch_plans.jsonl"
    if not plans_path.is_file():
        plans_path = supplementary_dir(repo, result.case_id, result.run_id) / "dispatch_plans.jsonl"
    with performance.measure("controller_chain_audit"):
        audit_controller_chain(run_dir, inputs, plans_path=plans_path)
    reserve_start = reserve_start_from_config(config)
    with performance.measure("physical_validation"):
        validation = validate_q4_run(
            result.intervals,
            executions,
            ledger,
            end_time=YEAR_END,
            actual_inputs=inputs,
            reserve_start=reserve_start,
        )
    if not validation["ok"]:
        raise InputError(
            "Q4 export independent validation failed: " + ";".join(validation["violations"][:8])
        )
    performance.switch_phase("export")
    with performance.measure("excel_export_and_readback"):
        template = (
            repo
            / "data"
            / "templates"
            / ("result4-2.xlsx" if result.case_id == "q4_2" else "result4-3.xlsx")
        )
        try:
            output.relative_to(run_dir)
        except ValueError:
            raise InputError("Q4 formal exported result must belong to its run directory") from None
        output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(template, output)
        wb = load_workbook(output)
        labels = [
            (sheet, wb[sheet]["B1"].value, wb[sheet]["EO1"].value)
            for sheet in ("计划购电量", "调整购电量")
            if sheet in wb.sheetnames
        ]
        try:
            mappings, expected = fill_workbook(wb, result, ledger)
            wb.save(output)
        finally:
            wb.close()
        readback = load_workbook(output, data_only=True)
        try:
            for sheet, row, col, value, tolerance in expected:
                actual = readback[sheet].cell(row, col).value
                mismatch = (
                    actual != value
                    if tolerance is None
                    else actual is None or abs(float(actual) - value) > tolerance
                )
                if mismatch:
                    raise InputError(f"Q4 readback mismatch: {sheet}!{get_column_letter(col)}{row}")
            for sheet, first, last in labels:
                if readback[sheet]["B1"].value != first or readback[sheet]["EO1"].value != last:
                    raise InputError("Q4 output modified template interval labels")
            if readback["充放电量"].max_row != 2005:
                raise InputError("Q4 charge sheet must have 334 six-row blocks")
            for sheet in ("充放电量", "紧急购电量"):
                if any("⁝" in str(cell.value) for row in readback[sheet] for cell in row):
                    raise InputError("Q4 output still has ellipsis rows")
        finally:
            readback.close()
    decisions = load_decisions(repo)
    decision = decisions["D-TIME-TEMPLATE-EXPORT-Q4"]
    decision_snapshot = {"D_TIME_TEMPLATE_EXPORT_Q4": decision}
    if reserve_start is not None:
        decision_snapshot["D_TERMINAL_RESERVE_Q4"] = decisions["D-TERMINAL-RESERVE-Q4"]
    if config["model_version"] == BLEND_MODEL_VERSION:
        decision_snapshot.update(
            {
                decision_id: decisions[decision_id.replace("_", "-")]
                for decision_id in BLEND_DECISIONS
            }
        )
    if (
        config["model_version"] == BIAS_MODEL_VERSION
        or config.get("solver_method", "scipy") != "scipy"
    ):
        decision_snapshot["D_OPTIMIZATION_Q4"] = decisions["D-OPTIMIZATION-Q4"]
    atomic_json(
        run_dir / "export_manifest.json",
        {
            "schema_version": 1,
            "case_id": result.case_id,
            "run_id": result.run_id,
            "source_result_sha256": sha256_file(domain_path),
            "output_sha256": sha256_file(output),
            "template_file": str(template.relative_to(repo)),
            "template_sha256": sha256_file(template),
            "decision_snapshot": decision_snapshot,
            "model_version": config["model_version"],
            "evidence_manifest_sha256": sha256_file(run_dir / "evidence_manifest.json"),
            "independent_validation": validation,
            "source_hashes": inputs.source_hashes,
            "readback_ok": True,
            "checked_cells": len(expected),
            "source_slot_mapping": mappings,
            "interval_label_conflicts": [
                {
                    "sheet": sheet,
                    "first_label": first,
                    "last_label": last,
                    "interpretation": "explicit column order B:EO; natural-day slots 0..143",
                }
                for sheet, first, last in labels
            ],
            "cost_columns": {
                "计划购电量!EQ": "planned_cost_cny",
                "调整购电量!EQ": "adjustment_cost_cny",
            },
            "total_cost_cny": ledger.costs().total_cost_cny,
        },
    )
    manifest_path = run_dir / "manifest.json"
    manifest = json.loads(manifest_path.read_text())
    try:
        relative_output = str(output.relative_to(run_dir))
    except ValueError:
        raise InputError("Q4 formal exported result must belong to its run directory") from None
    manifest["result_files"] = {output.name: relative_output}
    manifest["result_sha256"] = {output.name: sha256_file(output)}
    atomic_json(manifest_path, manifest)
    return output


def export_q4(repo: Path, result: CaseResult, output: Path) -> Path:
    performance = Performance()
    performance.switch_phase("validation")
    exported = _export_q4(repo, result, output, performance)
    timing = performance.summary()
    run = repo / "outputs" / "runs" / result.case_id / result.run_id
    run_timing = run / "performance_summary.json"
    timing.update(
        {
            "case_id": result.case_id,
            "run_id": result.run_id,
            "measurement_scope": "export_process_attempt",
            "export_included": True,
        }
    )
    if run_timing.is_file():
        previous = json.loads(run_timing.read_text())
        timing["run_measurement_sha256"] = sha256_file(run_timing)
        timing["run_measurement"] = previous
        timing["export_process_seconds"] = timing["end_to_end_seconds"]
        for name in ("simulation_seconds", "validation_seconds", "report_seconds"):
            timing[name] += previous[name]
        timing["end_to_end_seconds"] += previous["end_to_end_seconds"]
        timing["measurement_scope"] = (
            "sum_of_active_run_and_export_attempts_excluding_interprocess_wait"
        )
    atomic_json(output.with_suffix(".performance_summary.json"), timing)
    return exported
