"""Approved Q3 full-year delivery, preserving audited diagnostic sources."""

from __future__ import annotations

import datetime as dt
import json
import math
import re
import shutil
from copy import copy
from dataclasses import asdict, replace
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..approvals import load_decisions, require_approved_decisions
from ..artifacts import (
    build_manifest,
    ensure_run_id_available,
    file_digest,
    input_hashes,
    verify_imported_inputs,
    write_json,
    write_manifest,
)
from ..schemas import InputError, PendingDecisionError
from ..timekeys import parse_time_label
from .contracts import ENERGY_ABS_TOL_KWH, CaseResult, TimeGrid
from .q3 import REQUIRED_DECISIONS
from .q3_reporting import COST_KEYS, emergency_segments, settlement_summary, write_q3_paper_assets
from .q3_runtime_inputs import (
    Q3_EVALUATION_END,
    Q3_EVALUATION_START,
    inclusive_days,
    load_q3_runtime_inputs,
)
from .q3_sidecars import load_q3_sidecar
from .q3_terminal_reserve import validate_q3_terminal_reserve
from .result_io import load_case_result, save_case_result
from .validation import validate_complete_run

MAPPING_VERSION = "Q3-COLUMN-ORDER-v1"
SIDECAR_NAMES = ("forecast_provenance.jsonl", "plan_versions.jsonl", "settlement_ledger.jsonl")


def require_q3_export_approval(repo: Path) -> dict:
    require_approved_decisions(repo, (*REQUIRED_DECISIONS, "D_TIME_TEMPLATE_EXPORT"))
    decision = load_decisions(repo)["D-TIME-TEMPLATE-EXPORT"]
    if (
        decision.get("q3_scope_cases") != ["q3"]
        or decision.get("q3_mapping_version") != MAPPING_VERSION
        or not str(decision.get("q3_approval_source", "")).strip()
        or not str(decision.get("q3_choice", "")).strip()
    ):
        raise PendingDecisionError(
            ["D_TIME_TEMPLATE_EXPORT"], "Q1 approval does not cover the Q3 mapping"
        )
    return decision


def validate_annual_result(result: CaseResult) -> dict:
    days = inclusive_days(Q3_EVALUATION_START, Q3_EVALUATION_END)
    if result.case_id != "q3" or result.status != "success" or result.is_synthetic:
        raise InputError("Q3 delivery requires a successful nonsynthetic Q3 result")
    if len(result.intervals) != 48096 or {i.day for i in result.intervals} != set(days):
        raise InputError("Q3 delivery requires exactly 334 complete annual days")
    coverage = validate_complete_run(result.intervals, days)
    reserve = validate_q3_terminal_reserve(result.intervals)
    if not coverage.ok or not reserve["ok"]:
        raise InputError("Q3 annual continuity or terminal reserve validation failed")
    if (
        abs(result.intervals[0].state_start.energy_kwh - 6000) > ENERGY_ABS_TOL_KWH
        or abs(result.intervals[-1].state_end.energy_kwh - 6000) > ENERGY_ABS_TOL_KWH
    ):
        raise InputError("Q3 realized initial/final SOC must be 6000")
    for i in result.intervals:
        c, d = i.action.charge_kwh, i.action.discharge_kwh
        spill = (
            i.adjusted_purchase_kwh
            + i.pv_used_kwh
            + d
            + i.emergency_purchase_kwh
            - i.load_kw / 6
            - c
        )
        curtail = i.pv_kw / 6 - i.pv_used_kwh
        if (
            spill < -ENERGY_ABS_TOL_KWH
            or spill > i.adjusted_purchase_kwh + ENERGY_ABS_TOL_KWH
            or curtail < -ENERGY_ABS_TOL_KWH
        ):
            raise InputError(f"Q3 bus balance/curtailment invalid at {i.day} slot {i.slot}")
        if i.emergency_purchase_kwh > ENERGY_ABS_TOL_KWH and (
            c > ENERGY_ABS_TOL_KWH
            or abs(spill) > ENERGY_ABS_TOL_KWH
            or curtail > ENERGY_ABS_TOL_KWH
        ):
            raise InputError("Q3 emergency cannot charge or coexist with surplus")
    return {
        "ok": True,
        "violations": [],
        "complete_run": asdict(coverage),
        "terminal_reserve": reserve,
        "interval_count": 48096,
        "state_start_kwh": result.intervals[0].state_start.energy_kwh,
        "state_end_kwh": result.intervals[-1].state_end.energy_kwh,
    }


def fill_q3_workbook(wb, result: CaseResult, daily: dict, days: tuple[dt.date, ...]) -> list[dict]:
    if wb.sheetnames != ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]:
        raise InputError("Q3 official template sheets/order differ")
    charge, emergency = wb["充放电量"], wb["紧急购电量"]
    styles = {
        s.title: [copy(s.cell(2, c)._style) for c in range(1, s.max_column + 1)]
        for s in (charge, emergency)
    }
    for sheet in (charge, emergency):
        for merge in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merge))
        sheet.delete_rows(2, sheet.max_row)
    mappings = []
    emergency_row = 2
    for index, day in enumerate(days):
        intervals = result.intervals[index * 144 : (index + 1) * 144]
        if len(intervals) != 144 or any(
            i.day != day or i.slot != k for k, i in enumerate(intervals)
        ):
            raise InputError("Q3 export day/slot coverage differs")
        date_value = wb["计划购电量"].cell(index + 2, 1).value
        for sheet_name in ("计划购电量", "调整购电量"):
            label = wb[sheet_name].cell(index + 2, 1).value
            if label is None or parse_time_label(label).base_date != day:
                raise InputError("Q3 template date order differs from source")
        for slot, interval in enumerate(intervals):
            wb["计划购电量"].cell(index + 2, slot + 2, interval.planned_purchase_kwh)
            wb["调整购电量"].cell(index + 2, slot + 2, interval.adjusted_purchase_kwh)
            period = TimeGrid().interval(day, slot)
            mappings.append(
                {
                    "day": str(day),
                    "slot": slot,
                    "interval_start": period.start.isoformat(),
                    "interval_end": period.end.isoformat(),
                    "original_label": wb["计划购电量"].cell(1, slot + 2).value,
                    "G0_cell": f"计划购电量!{get_column_letter(slot + 2)}{index + 2}",
                    "Gfinal_cell": f"调整购电量!{get_column_letter(slot + 2)}{index + 2}",
                    "charge_cell": f"充放电量!C{2 + index * 6 + slot // 24}",
                    "discharge_cell": f"充放电量!D{2 + index * 6 + slot // 24}",
                }
            )
        for sheet_name, quantity, cost in (
            (
                "计划购电量",
                math.fsum(i.planned_purchase_kwh for i in intervals),
                daily[day]["planned_cost_cny"],
            ),
            (
                "调整购电量",
                math.fsum(i.adjusted_purchase_kwh for i in intervals),
                daily[day]["adjustment_cost_cny"],
            ),
        ):
            wb[sheet_name].cell(index + 2, 146, quantity)
            wb[sheet_name].cell(index + 2, 147, cost)
        first_row = 2 + index * 6
        for block in range(6):
            row = first_row + block
            for col in range(1, 7):
                charge.cell(row, col)._style = copy(styles[charge.title][col - 1])
            charge.cell(row, 2, f"{block * 4}:00-{(block + 1) * 4}:00")
            part = intervals[block * 24 : (block + 1) * 24]
            charge.cell(row, 3, math.fsum(i.action.charge_kwh for i in part))
            charge.cell(row, 4, math.fsum(i.action.discharge_kwh for i in part))
        charge.cell(first_row, 1, date_value)
        charge.merge_cells(start_row=first_row, end_row=first_row + 5, start_column=1, end_column=1)
        for offset, label, state in (
            (0, "0:00", intervals[0].state_start.energy_kwh),
            (1, "24:00", intervals[-1].state_end.energy_kwh),
        ):
            charge.cell(first_row + offset, 5, label)
            charge.cell(first_row + offset, 6, state)
        groups = emergency_segments(intervals) or (
            {"label": None, "energy_kwh": 0.0, "first_slot": 0, "end_slot_exclusive": 0},
        )
        for group in groups:
            for col in range(1, 4):
                emergency.cell(emergency_row, col)._style = copy(styles[emergency.title][col - 1])
            emergency.cell(emergency_row, 1, date_value)
            emergency.cell(emergency_row, 2, group["label"])
            emergency.cell(emergency_row, 3, group["energy_kwh"])
            for slot in range(group["first_slot"], group["end_slot_exclusive"]):
                mappings[index * 144 + slot]["emergency_cell"] = f"紧急购电量!C{emergency_row}"
            emergency_row += 1
    return mappings


def readback_q3_workbook(
    template: Path, output: Path, result: CaseResult, daily: dict, days: tuple[dt.date, ...]
) -> dict:
    original, actual = load_workbook(template), load_workbook(output, data_only=True)
    checked = 0

    def check(sheet, row, column, expected, tolerance=None):
        nonlocal checked
        value = actual[sheet].cell(row, column).value
        if tolerance is None:
            mismatch = value != expected
        else:
            mismatch = (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(value)
                or abs(value - expected) > tolerance
            )
        if mismatch:
            raise InputError(
                f"Q3 independent readback mismatch: {sheet}!{get_column_letter(column)}{row}"
            )
        checked += 1

    try:
        if actual.sheetnames != original.sheetnames:
            raise InputError("Q3 readback sheet order differs")
        for sheet in ("计划购电量", "调整购电量"):
            if actual[sheet].max_row != len(days) + 1 or actual[sheet].max_column != 147:
                raise InputError("Q3 readback purchase table dimensions differ")
            for column in range(1, 148):
                check(sheet, 1, column, original[sheet].cell(1, column).value)
        emergency_row = 2
        for index, day in enumerate(days):
            intervals = result.intervals[index * 144 : (index + 1) * 144]
            for sheet, field, cost_key in (
                ("计划购电量", "planned_purchase_kwh", "planned_cost_cny"),
                ("调整购电量", "adjusted_purchase_kwh", "adjustment_cost_cny"),
            ):
                check(sheet, index + 2, 1, original[sheet].cell(index + 2, 1).value)
                for slot, i in enumerate(intervals):
                    check(sheet, index + 2, slot + 2, getattr(i, field), ENERGY_ABS_TOL_KWH)
                check(
                    sheet,
                    index + 2,
                    146,
                    math.fsum(getattr(i, field) for i in intervals),
                    ENERGY_ABS_TOL_KWH,
                )
                check(sheet, index + 2, 147, daily[day][cost_key], 1e-5)
            first = 2 + 6 * index
            date_value = original["计划购电量"].cell(index + 2, 1).value
            check("充放电量", first, 1, date_value)
            if f"A{first}:A{first + 5}" not in actual["充放电量"].merged_cells:
                raise InputError("Q3 readback missing six-row date merge")
            for block in range(6):
                part = intervals[block * 24 : (block + 1) * 24]
                check("充放电量", first + block, 2, f"{block * 4}:00-{(block + 1) * 4}:00")
                for column, field in ((3, "charge_kwh"), (4, "discharge_kwh")):
                    check(
                        "充放电量",
                        first + block,
                        column,
                        math.fsum(getattr(i.action, field) for i in part),
                        ENERGY_ABS_TOL_KWH,
                    )
            for offset, label, state in (
                (0, "0:00", intervals[0].state_start.energy_kwh),
                (1, "24:00", intervals[-1].state_end.energy_kwh),
            ):
                check("充放电量", first + offset, 5, label)
                check("充放电量", first + offset, 6, state, ENERGY_ABS_TOL_KWH)
            for offset in range(2, 6):
                for col in (5, 6):
                    check("充放电量", first + offset, col, None)
            for group in emergency_segments(intervals) or ({"label": None, "energy_kwh": 0.0},):
                check("紧急购电量", emergency_row, 1, date_value)
                check("紧急购电量", emergency_row, 2, group["label"])
                check("紧急购电量", emergency_row, 3, group["energy_kwh"], ENERGY_ABS_TOL_KWH)
                emergency_row += 1
        if (
            actual["充放电量"].max_row != 1 + 6 * len(days)
            or actual["紧急购电量"].max_row != emergency_row - 1
        ):
            raise InputError("Q3 readback still has omitted or extraneous rows")
        return {
            "ok": True,
            "checked_cells": checked,
            "purchase_intervals": len(result.intervals),
            "charge_rows": 6 * len(days),
            "emergency_rows": emergency_row - 2,
        }
    finally:
        original.close()
        actual.close()


def create_q3_delivery(
    repo: Path, *, source_run_id: str, run_id: str, audit_report: Path
) -> CaseResult:
    decision = require_q3_export_approval(repo)
    if source_run_id != decision.get("q3_formal_source_run_id") or run_id != decision.get(
        "q3_formal_run_id"
    ):
        raise PendingDecisionError(
            ["D_TIME_TEMPLATE_EXPORT"], "Q3 source/target differ from explicitly approved delivery"
        )
    if source_run_id == run_id or any(
        not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]*", v) for v in (source_run_id, run_id)
    ):
        raise InputError("Q3 source and target must be distinct simple run IDs")
    target = ensure_run_id_available(repo, "q3", run_id)
    source = repo / "outputs/runs/q3" / source_run_id
    if file_digest(audit_report) != decision.get("q3_audit_sha256"):
        raise InputError("Q3 audit receipt differs from approved SHA-256")
    receipt = json.loads(audit_report.read_text(encoding="utf-8"))
    manifest = json.loads((source / "manifest.json").read_text(encoding="utf-8"))
    if (
        receipt.get("run_id") != source_run_id
        or receipt.get("validation_ok") is not True
        or manifest.get("run_id") != source_run_id
        or manifest.get("case_id") != "q3"
        or manifest.get("status") != "success"
        or manifest.get("is_synthetic") is not False
        or manifest.get("diagnostic_only") is not True
        or manifest.get("validation_ok") is not True
    ):
        raise InputError("Q3 source is not the approved audited nonsynthetic diagnostic")
    if receipt.get("code_commit") != manifest.get("code_commit") or receipt.get(
        "source_hash"
    ) != manifest.get("source_hash"):
        raise InputError("Q3 audit source code binding differs")
    if verify_imported_inputs(repo) or manifest.get("input_hashes") != input_hashes(repo):
        raise InputError("Q3 inputs differ from verified source attachments")
    current = load_decisions(repo)
    for key in REQUIRED_DECISIONS:
        previous = manifest["config_snapshot"]["decisions.toml"]["decisions"][key]
        for field in ("status", "choice", "confirmed_by", "confirmed_at"):
            if previous.get(field) != current[key.replace("_", "-")].get(field):
                raise InputError(f"Q3 model approval drift: {key}.{field}")
    for name in ("manifest.json", "summary.json", "domain_result.json", *SIDECAR_NAMES):
        digest = file_digest(source / name)
        if digest != receipt["file_sha256"].get(name) or (
            name in SIDECAR_NAMES and digest != manifest["result_sha256"].get(name)
        ):
            raise InputError(f"Q3 audited file hash changed: {name}")
    result = load_case_result(source / "domain_result.json")
    if result.run_id != source_run_id:
        raise InputError("Q3 source domain identity differs")
    validation = validate_annual_result(result)
    runtime = load_q3_runtime_inputs(
        attachment1_path=repo / "data/raw/附件1.xlsx",
        attachment2_path=repo / "data/raw/附件2.xlsx",
        attachment3_path=repo / "data/raw/附件3.xlsx",
    )
    actuals = runtime.actuals_for(inclusive_days(Q3_EVALUATION_START, Q3_EVALUATION_END))
    if any(
        (i.day, i.slot, i.load_kw, i.pv_kw) != (a.day, a.slot, a.load_kw, a.pv_kw)
        for i, a in zip(result.intervals, actuals, strict=True)
    ):
        raise InputError("Q3 persisted actuals differ from original inputs")
    rows = load_q3_sidecar(source / "settlement_ledger.jsonl")
    daily, _ = settlement_summary(rows)
    summary = json.loads((source / "summary.json").read_text(encoding="utf-8"))
    costs = {
        key: math.fsum(row["cost_cny"] for row in rows if row["record_type"] == kind)
        for kind, key in COST_KEYS.items()
    }
    costs["total_cost_cny"] = math.fsum(costs.values())
    for key, value in costs.items():
        if abs(value - result.metadata[key]) > 1e-5 or abs(value - summary[key]) > 1e-5:
            raise InputError("Q3 actual settlement does not reconcile to summary")
    validation.update(
        actual_input_match=True, costs_cny=costs, source_audit_sha256=file_digest(audit_report)
    )
    template = repo / "data/templates/result3.xlsx"
    template_hash = file_digest(template)
    target.mkdir(parents=True, exist_ok=True)
    provenance = {
        "operation": "user-approved formal delivery from audited evidence; no MPC rerun",
        "source_run_id": source_run_id,
        "source_manifest_sha256": file_digest(source / "manifest.json"),
        "source_domain_sha256": file_digest(source / "domain_result.json"),
        "source_audit_sha256": file_digest(audit_report),
        "trajectory_source": manifest["trajectory_source"],
        "serialization_commit": manifest["code_commit"],
        "approval_source": decision["q3_approval_source"],
    }
    for name in SIDECAR_NAMES:
        shutil.copy2(source / name, target / name)
        if file_digest(target / name) != receipt["file_sha256"][name]:
            raise InputError("Q3 sidecar copy failed hash validation")
    metadata = {**result.metadata, "formal_delivery_source": provenance}
    delivered = replace(
        result,
        run_id=run_id,
        metadata=metadata,
        result_files=tuple(
            target / name for name in ("domain_result.json", *SIDECAR_NAMES, "results/result3.xlsx")
        ),
    )
    save_case_result(target / "domain_result.json", delivered)
    if load_case_result(target / "domain_result.json") != delivered:
        raise InputError("Q3 formal domain failed exact roundtrip")
    summary.update(run_id=run_id, formal_delivery_source=provenance, diagnostic_only=False)
    write_json(target / "summary.json", summary)
    output = target / "results/result3.xlsx"
    output.parent.mkdir()
    wb = load_workbook(template)
    days = inclusive_days(Q3_EVALUATION_START, Q3_EVALUATION_END)
    try:
        mappings = fill_q3_workbook(wb, delivered, daily, days)
        wb.save(output)
    finally:
        wb.close()
    readback = readback_q3_workbook(template, output, delivered, daily, days)
    if file_digest(template) != template_hash:
        raise InputError("Q3 original template changed")
    validation["excel_readback"] = readback
    write_json(target / "validation.json", validation)
    write_json(
        target / "export_manifest.json",
        {
            "case_id": "q3",
            "run_id": run_id,
            "source_result_sha256": file_digest(target / "domain_result.json"),
            "output_sha256": file_digest(output),
            "template_file": "data/templates/result3.xlsx",
            "template_sha256": template_hash,
            "decision_snapshot": {"D_TIME_TEMPLATE_EXPORT": decision},
            "mapping_version": MAPPING_VERSION,
            "source_slot_mapping": mappings,
            "readback_ok": True,
            "readback": readback,
            "cost_columns": {
                "计划购电量!EQ": "planned_cost_cny",
                "调整购电量!EQ": "adjustment_cost_cny",
            },
            "costs_cny": costs,
            "provenance": provenance,
        },
    )
    asset_manifest = write_q3_paper_assets(delivered, rows, target / "paper_assets")
    names = (
        "domain_result.json",
        "summary.json",
        "validation.json",
        "export_manifest.json",
        *SIDECAR_NAMES,
        "results/result3.xlsx",
        "paper_assets/asset_manifest.json",
    )
    final_manifest = build_manifest(
        repo,
        run_id=run_id,
        case_id="q3",
        command=[
            "python",
            "scripts/export_q3_run.py",
            "--source-run-id",
            source_run_id,
            "--run-id",
            run_id,
        ],
        status="success",
        is_synthetic=False,
        model_status="formal_delivery_from_audited_trajectory",
        result_files={Path(n).name: n for n in names},
        result_sha256={Path(n).name: file_digest(target / n) for n in names},
        verify_inputs=True,
    )
    final_manifest.update(
        validation_ok=True,
        diagnostic_only=False,
        formal_delivery=True,
        artifact_regeneration_only=True,
        formal_delivery_source=provenance,
        paper_asset_manifest="paper_assets/asset_manifest.json",
        paper_asset_sha256=file_digest(asset_manifest),
    )
    write_manifest(target, final_manifest)
    return delivered


def publish_q3_paper_assets(repo: Path, run_id: str, output_dir: Path | None = None) -> Path:
    """Publish only the explicitly selected Q3 delivery's hash-checked assets."""
    import tomllib

    from ..checks import check_selected_run

    decision = require_q3_export_approval(repo)
    selected = tomllib.loads((repo / "configs/selected_runs.toml").read_text(encoding="utf-8"))
    if (
        run_id != decision.get("q3_formal_run_id")
        or selected.get("runs", {}).get("q3") != run_id
        or selected.get("case_selections", {}).get("q3", {}).get("selection_status") != "approved"
    ):
        raise InputError("Q3 paper assets require the explicitly approved case selection")
    blockers = check_selected_run(repo, "q3", run_id)
    if blockers:
        raise InputError("Q3 selected delivery invalid: " + "; ".join(blockers))
    source = repo / "outputs/runs/q3" / run_id
    manifest = json.loads((source / "manifest.json").read_text(encoding="utf-8"))
    asset_path = source / "paper_assets/asset_manifest.json"
    if (
        manifest.get("diagnostic_only") is not False
        or manifest.get("formal_delivery") is not True
        or file_digest(asset_path) != manifest.get("paper_asset_sha256")
    ):
        raise InputError("Q3 paper source must be an audited formal delivery")
    assets = json.loads(asset_path.read_text(encoding="utf-8"))
    if (
        assets.get("run_id") != run_id
        or assets.get("source_result_sha256") != file_digest(source / "domain_result.json")
        or assets.get("is_synthetic") is not False
    ):
        raise InputError("Q3 asset provenance differs from the selected domain result")
    for name, digest in assets["files_sha256"].items():
        if Path(name).name != name or file_digest(source / "paper_assets" / name) != digest:
            raise InputError("Q3 paper asset hash mismatch")
    destination = output_dir or repo / "paper/generated"
    destination.mkdir(parents=True, exist_ok=True)
    for name in ("q3_result_tables.tex", "q3_dispatch_days.pdf", "q3_dispatch_days.png"):
        target = destination / name
        if target.exists() and file_digest(target) != assets["files_sha256"][name]:
            raise InputError("refusing to overwrite Q3 paper assets from a different source")
        shutil.copy2(source / "paper_assets" / name, target)
    return destination / "q3_result_tables.tex"


def export_q3_case_result(repo: Path, result: CaseResult, output: Path) -> Path:
    """Shared export API: write only an exact, explicitly approved saved delivery."""
    decision = require_q3_export_approval(repo)
    if result.run_id != decision.get("q3_formal_run_id"):
        raise PendingDecisionError(
            ["D_TIME_TEMPLATE_EXPORT"], "Q3 result is not the approved formal delivery"
        )
    validate_annual_result(result)
    source = repo / "outputs/runs/q3" / result.run_id
    if load_case_result(source / "domain_result.json") != result:
        raise InputError("Q3 export result differs from saved formal source")
    try:
        output.resolve().relative_to(source.resolve())
    except ValueError:
        raise InputError("Q3 export must belong to its approved delivery directory") from None
    if output.exists():
        raise InputError("Q3 export output already exists")
    daily, _ = settlement_summary(load_q3_sidecar(source / "settlement_ledger.jsonl"))
    template = repo / "data/templates/result3.xlsx"
    template_hash = file_digest(template)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(template)
    days = inclusive_days(Q3_EVALUATION_START, Q3_EVALUATION_END)
    try:
        fill_q3_workbook(wb, result, daily, days)
        wb.save(output)
    finally:
        wb.close()
    readback_q3_workbook(template, output, result, daily, days)
    if file_digest(template) != template_hash:
        raise InputError("Q3 original template changed")
    return output
